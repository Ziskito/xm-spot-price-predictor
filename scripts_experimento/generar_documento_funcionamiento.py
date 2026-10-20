# -*- coding: utf-8 -*-
"""
Documento de referencia "Augur: funcionamiento del sistema" (docs/Funcionamiento_Augur.docx), 19-oct-2026.
Explica el sistema tal como corre hoy: datos y calendario de publicacion, las 66 variables, los modelos finales, el
motor de decision, las metricas de calidad (que son y por que se usan), los resultados, la comparacion con otros
trabajos, la operacion diaria y las limitaciones. Las cifras salen de informe_final_assets/cifras.json (las mismas
del informe final); las figuras nuevas se dibujan en funcionamiento_assets/.
Uso: python scripts_experimento/generar_documento_funcionamiento.py [--hasta N]   (N = ultima seccion a incluir)
"""
import json
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.section import WD_SECTION  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parent
sys.path.insert(0, str(AQUI))
from ecuaciones_omml import omml  # noqa: E402

DESTINO = RAIZ / "docs" / "Funcionamiento_Augur.docx"
FIG_INF = AQUI / "informe_final_assets"
FIG = AQUI / "funcionamiento_assets"
FIG.mkdir(exist_ok=True)
C = json.loads((FIG_INF / "cifras.json").read_text(encoding="utf-8"))
HASTA = int(sys.argv[sys.argv.index("--hasta") + 1]) if "--hasta" in sys.argv else 99

AZUL, GRIS, TINTA = RGBColor(0x1F, 0x4E, 0x79), RGBColor(0x59, 0x59, 0x59), RGBColor(0x1A, 0x1A, 0x1A)
FONDO_ENC, FONDO_NOTA = "1F4E79", "EAF1F8"


# ------------------------------------------------------------------------------------------------ formato de numeros
def n(v, d=1):
    s = f"{abs(v):,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return ("−" if v < 0 else "") + s


def pc(v, d=1):
    return n(v, d) + " %"


# ------------------------------------------------------------------------------------------------ documento base
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
sec.left_margin = sec.right_margin = Cm(2.3)
sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.2)

st = doc.styles["Normal"]
st.font.name, st.font.size, st.font.color.rgb = "Calibri", Pt(11), TINTA
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.paragraph_format.space_after, st.paragraph_format.line_spacing = Pt(6), 1.12
for nivel, tam, antes in ((1, 17, 18), (2, 13, 12), (3, 11.5, 9)):
    h = doc.styles[f"Heading {nivel}"]
    h.font.name, h.font.size, h.font.bold, h.font.color.rgb = "Calibri", Pt(tam), True, AZUL
    rf = h.element.rPr.find(qn("w:rFonts"))
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf is not None and rf.get(qn(a)):
            del rf.attrib[qn(a)]
    h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(antes), Pt(4)
    h.paragraph_format.keep_with_next = True

ANCHO = sec.page_width - sec.left_margin - sec.right_margin


def sombra(celda, color):
    tcpr = celda._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"), s.set(qn("w:color"), "auto"), s.set(qn("w:fill"), color)
    tcpr.append(s)


def runs(par, texto, tam=None, color=None):
    """**negrita** y *cursiva* dentro del texto."""
    import re
    for parte in re.split(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*))", texto):
        if not parte:
            continue
        if parte.startswith("**"):
            r = par.add_run(parte[2:-2]); r.bold = True
        elif parte.startswith("*"):
            r = par.add_run(parte[1:-1]); r.italic = True
        else:
            r = par.add_run(parte)
        if tam:
            r.font.size = Pt(tam)
        if color:
            r.font.color.rgb = color
    return par


def P(texto, estilo=None, alin=None):
    p = doc.add_paragraph(style=estilo)
    runs(p, texto)
    if alin:
        p.alignment = alin
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def V(texto, nivel=0):
    p = doc.add_paragraph(style="List Bullet" if nivel == 0 else "List Bullet 2")
    runs(p, texto)
    p.paragraph_format.space_after = Pt(3)
    return p


def H(texto, nivel=1):
    return doc.add_heading(texto, level=nivel)


def nota(texto, titulo="Por qué importa"):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    sombra(c, FONDO_NOTA)
    p = c.paragraphs[0]
    runs(p, f"**{titulo}.** {texto}", tam=10)
    p.paragraph_format.space_after = Pt(2)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def tabla(enc, filas, anchos=None, tam=9, titulo=None, fuente=None):
    if titulo:
        p = doc.add_paragraph()
        runs(p, f"**{titulo}**", tam=10, color=AZUL)
        p.paragraph_format.space_after, p.paragraph_format.keep_with_next = Pt(3), True
    t = doc.add_table(rows=1, cols=len(enc))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, e in enumerate(enc):
        c = t.rows[0].cells[j]
        sombra(c, FONDO_ENC)
        runs(c.paragraphs[0], f"**{e}**", tam=tam, color=RGBColor(0xFF, 0xFF, 0xFF))
    for i, f in enumerate(filas):
        cs = t.add_row().cells
        for j, v in enumerate(f):
            runs(cs[j].paragraphs[0], str(v), tam=tam)
            if i % 2 == 1:
                sombra(cs[j], "F3F6FA")
    for fila in t.rows:
        for c in fila.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.0
    if anchos:
        tot = sum(anchos)
        for fila in t.rows:
            for j, c in enumerate(fila.cells):
                c.width = int(ANCHO * anchos[j] / tot)
    # encabezado repetido en cada pagina
    trpr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trpr.append(th)
    if fuente:
        p = doc.add_paragraph()
        runs(p, fuente, tam=8.5, color=GRIS)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def ec(latex, explicacion=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p._p.append(omml(latex))
    p.paragraph_format.space_after = Pt(4)
    if explicacion:
        P(explicacion)


def figura(ruta, ancho_cm, pie):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ruta), width=Cm(ancho_cm))
    p.paragraph_format.keep_with_next = True
    q = doc.add_paragraph()
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    runs(q, pie, tam=9, color=GRIS)


def salto():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ------------------------------------------------------------------------------------------------ figuras nuevas
def fig_calendario():
    ruta = FIG / "calendario_publicacion.png"
    fuentes = [
        ("Precio de bolsa, demanda y generación (API XM)", -4.6, -3, "#c0392b", "hasta 23:00 de D-4"),
        ("Embalses y aportes (API XM)", -4.6, -1, "#2e86c1", "valor diario de D-2"),
        ("Ofertas de precio de los generadores (cierre 08:00)", -1, -1 + 8 / 24, "#7f8c8d", "no públicas a tiempo"),
        ("Predespacho ideal de XM para D", -1, -1 + 10.5 / 24, "#27ae60", "D-1 ~10:30"),
        ("Despacho programado (dMAR, dDEM) para D", -1, -1 + 14.5 / 24, "#16a085", "D-1 ~14:30"),
    ]
    fig, ax = plt.subplots(figsize=(10, 3.4), dpi=170)
    for i, (nom, ini, fin, col, txt) in enumerate(fuentes[::-1]):
        if ini == fin or nom.startswith("Ofertas") or nom.startswith("Predespacho") or nom.startswith("Despacho"):
            ax.plot([fin], [i], marker="o", ms=9, color=col)
            ax.text(1.08, i, txt, va="center", fontsize=8.5, color=col)
        else:
            ax.barh(i, fin - ini, left=ini, height=0.45, color=col, alpha=0.85)
            ax.text(ini + 0.05, i + 0.33, "datos conocidos", fontsize=7.5, color=col)
            ax.text(1.08, i, txt, va="center", fontsize=8.5, color=col)
    ax.set_yticks(range(len(fuentes)))
    ax.set_yticklabels([f[0] for f in fuentes[::-1]], fontsize=9)
    ax.axvline(0, color="black", lw=1.6)
    ax.axvspan(0, 1, color="#f5b041", alpha=0.18)
    ax.text(0.5, len(fuentes) - 0.35, "Día pronosticado D\n(01:00 a 00:00)", ha="center", va="center", fontsize=8.5)
    ax.text(0.03, -0.7, "Corte: 00:00 de D", fontsize=8.5, fontweight="bold")
    ax.set_xlim(-4.6, 2.3)
    ax.set_ylim(-1, len(fuentes))
    ax.set_xticks([-4, -3, -2, -1, 0, 1])
    ax.set_xticklabels(["D-4", "D-3", "D-2", "D-1", "D", "D+1"])
    for s in ("top", "right", "left"):
        ax.spines[s].set_visible(False)
    ax.tick_params(axis="y", length=0)
    fig.tight_layout()
    fig.savefig(ruta)
    plt.close(fig)
    return ruta


def fig_comparacion():
    ruta = FIG / "comparacion_mae.png"
    orden = [(0, "Persistencia"), (1, "XGBoost"), (2, "ARX+GARCH"), (3, "N-BEATSx"), (4, "N-HiTS"), (5, "GARCH-ged"),
             (6, "Ensamble de 6"), (7, "Predespacho XM"), (10, "Augur (definitivo)")]
    vals, etq = [C["tabla24"][i]["MAE"] for i, _ in orden], [e for _, e in orden]
    fig, ax = plt.subplots(figsize=(9, 3.6), dpi=170)
    cols = ["#95a5a6"] * 7 + ["#27ae60", "#1f4e79"]
    ax.barh(range(len(vals))[::-1], vals, color=cols)
    for i, v in enumerate(vals):
        ax.text(v + 1.5, len(vals) - 1 - i, n(v), va="center", fontsize=9)
    ax.set_yticks(range(len(vals))[::-1]); ax.set_yticklabels(etq, fontsize=9)
    ax.set_xlabel("MAE en 2026 (COP/kWh), mismas 4.848 horas, solo información pública al corte", fontsize=9)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.set_xlim(0, max(vals) * 1.12)
    fig.tight_layout(); fig.savefig(ruta); plt.close(fig)
    return ruta


# ================================================================================================== contenido
K = C["contrato24"]
T24 = C["tabla24"]
T72 = {f["tramo"]: f for f in C["tabla72"]}
OP = C["operacion"]
mejora_pers = 100 * (1 - K["MAE"] / K["MAE_pers"])
mejora_pre = 100 * (1 - K["MAE"] / K["MAE_pre"])


def s0_portada():
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("Augur"); r.bold = True; r.font.size = Pt(40); r.font.color.rgb = AZUL
    p = doc.add_paragraph()
    r = p.add_run("Funcionamiento del sistema"); r.font.size = Pt(22); r.font.color.rgb = TINTA
    p = doc.add_paragraph()
    runs(p, "Pronóstico del precio de bolsa de energía en Colombia y motor de decisión: datos, variables, modelos, "
            "métricas de calidad, resultados y operación", tam=12.5, color=GRIS)
    for _ in range(9):
        doc.add_paragraph()
    for linea in ("**Juan David Barceló Barraza · Rafael Andrés Dede Perdomo**",
                  "Ingeniería Electrónica · Universidad del Norte",
                  "Documento de referencia técnica · versión del 19 de octubre de 2026",
                  "Dashboard: augurpredictor.streamlit.app"):
        p = doc.add_paragraph(); runs(p, linea, tam=11, color=GRIS); p.paragraph_format.space_after = Pt(2)
    salto()
    H("Contenido", 1)
    for i, t in enumerate(["Resumen: qué hace Augur y cifras clave", "El mercado y qué se pronostica",
                           "Datos e información disponible al corte", "Variables de entrada (66)",
                           "Modelos finales", "Motor de decisión", "Métricas de calidad: qué son y por qué se usan",
                           "Cómo se validó", "Resultados", "Comparación con otros trabajos", "Operación diaria",
                           "Limitaciones conocidas", "Glosario"], 1):
        p = doc.add_paragraph(); runs(p, f"{i}.  {t}", tam=11.5); p.paragraph_format.space_after = Pt(3)
    salto()


def s1_resumen():
    H("1. Resumen: qué hace Augur y cifras clave")
    P("Augur pronostica, hora por hora, el **precio de bolsa nacional** del mercado mayorista colombiano para el día "
      "siguiente (24 h) y para los tres días siguientes (72 h). Para cada hora entrega un valor central (la mediana, q50), "
      "una banda de incertidumbre del 80 % (q10 a q90) y la probabilidad de que ocurra un **pico** o un **desplome** de "
      "precio. Sobre ese pronóstico, un **motor de decisión** traduce los números en una recomendación por hora para dos "
      "tipos de usuario: el generador (vender, retener o esperar) y el comercializador (comprar, evitar la compra o "
      "esperar). Todo se muestra en un dashboard web público que se actualiza solo tres veces al día.")
    P("La idea central del diseño definitivo es sencilla: **no se pronostica el precio desde cero, se corrige el "
      "predespacho ideal que XM publica el día anterior**. El predespacho ya contiene las ofertas de los generadores y la "
      "demanda esperada, que son las que fijan el precio; lo que Augur aprende es en qué condiciones el predespacho se "
      "equivoca y cuánto. Tres modelos CatBoost estiman esa corrección y un combinador causal mezcla el resultado con el "
      "predespacho.")
    P("Regla que atraviesa todo el proyecto: **solo se usa información pública que ya existe en el momento de "
      "pronosticar** (el corte, a las 00:00 del día pronosticado). XM publica el precio con unos tres días de rezago, así "
      "que el modelo nunca ve el precio de ayer. Esta restricción hace la tarea más difícil, pero es la que enfrentaría "
      "un usuario real con datos públicos.")
    filas = [
        ["Error absoluto medio (MAE), 24 h, 2026", f"{n(K['MAE'])} COP/kWh", f"sobre {K['horas']:,} horas (ene-ago 2026)".replace(",", ".")],
        ["Error porcentual (MAPE), 24 h", pc(K["MAPE"]), "el precio medio del periodo fue ~395 COP/kWh"],
        ["Mejora frente a la persistencia", f"{pc(mejora_pers, 0)} menos error", f"persistencia: {n(K['MAE_pers'])} COP/kWh; DM t = {n(K['DM_vs_pers'][0])}"],
        ["Mejora frente al predespacho de XM", f"{pc(mejora_pre, 0)} menos error", f"predespacho: {n(K['MAE_pre'])} COP/kWh; DM t = {n(K['DM_vs_pre'][0])}"],
        ["MAE a 72 h (pasos 1 a 72)", f"{n(T72['1-72 h']['MAE'])} COP/kWh", f"25-48 h: {n(T72['25-48 h']['MAE'])} · 49-72 h: {n(T72['49-72 h']['MAE'])}"],
        ["Cobertura de la banda del 80 %", pc(K["cobertura"]), "objetivo 80 %; ancho medio " + n(K["ancho"], 0) + " COP/kWh"],
        ["Detección de picos (AUC)", n(C["eventos"][5]["AUC"], 2), "desplomes: AUC " + n(C["eventos"][2]["AUC"], 2)],
        ["Operación en vivo (5-ago a 11-oct)", f"MAE {n(OP['MAE'])} COP/kWh", f"predespacho {n(OP['MAE_pre'])}; incluye el cambio de reglas de sep-2026"],
        ["Validación histórica (5 años, 2020-2025)", "gana al predespacho en 5 de 5", "y a la versión simple en 3 de 5 (empata en 2)"],
    ]
    tabla(["Indicador", "Valor", "Contexto"], filas, anchos=[4.2, 2.6, 5.2], titulo="Tabla 1. Cifras clave del sistema",
          fuente="COP/kWh: pesos colombianos por kilovatio-hora. DM: estadístico de Diebold-Mariano (sección 7); valores "
                 "por encima de 2 indican una diferencia significativa.")
    figura(FIG_INF / "diagrama_bloques.png", 15.5, "Figura 1. Arquitectura de Augur: datos (OE1), pronóstico (OE2) y "
                                                    "decisión y validación (OE3/OE4).")


def s2_mercado():
    H("2. El mercado y qué se pronostica")
    H("2.1 Cómo se forma el precio de bolsa", 2)
    P("En el mercado mayorista colombiano cada generador ofrece, antes de las 08:00 del día anterior, **un único precio "
      "para las 24 horas** del día siguiente y su disponibilidad hora a hora. Con esas ofertas, XM (el operador del "
      "sistema y administrador del mercado) arma el despacho: ordena las plantas de la más barata a la más cara hasta "
      "cubrir la demanda. Hay dos despachos que importan para Augur:")
    V("**Predespacho ideal (antes):** XM lo publica el día anterior hacia las 10:30. Es el despacho sin restricciones de "
      "la red con la demanda pronosticada; su costo marginal por hora es la mejor estimación pública del precio.")
    V("**Despacho ideal (después):** se calcula ex post con la demanda real. El precio de bolsa de cada hora es el "
      "**máximo precio de oferta (MPO)** de las plantas despachadas en ese despacho ideal, más un componente adicional "
      "(ΔI) que reconoce costos de arranque y parada.")
    P("Por eso el predespacho y el precio real difieren: la demanda real no es la pronosticada, la hidrología cambia y "
      "las plantas fallan. Esa diferencia es exactamente lo que Augur aprende a corregir.")
    H("2.2 Versiones del precio y cuál usamos", 2)
    tabla(["Versión", "Cuándo se publica", "Qué es"], [
        ["TX1", "hacia D+2", "primera liquidación del día; es la que alimenta la API pública"],
        ["TX2", "hacia D+5", "corrige la TX1 con medidas revisadas"],
        ["TXR", "días 1 a 5 del mes siguiente", "reliquidación mensual"],
        ["TXF", "hacia el día 10 del mes siguiente", "versión final"],
    ], anchos=[1.4, 3.6, 7], titulo="Tabla 2. Versiones de liquidación del precio (XM)")
    P("Augur pronostica el **precio de bolsa nacional (PBN)** tal como lo publica la API de XM. Como la TX1 sale hacia "
      "D+2 y el sistema público la refleja con algo más de rezago, **al corte del día D el último precio conocido es el "
      "de las 23:00 de D-4**. No hay forma pública de tener el precio de ayer.")
    H("2.3 Reglas que cambiaron en 2026", 2)
    V("**Precio de transacciones en bolsa (PTB) frente a PBN.** Desde la Resolución CREG 101 066 de 2024 hay tres "
      "precios de escasez y el precio con que se liquidan las transacciones (PTB) puede quedar por debajo del PBN. En "
      "2026 eso pasó en 360 horas. Augur pronostica el PBN; quien liquide al PTB debe tenerlo en cuenta.")
    V("**Resolución CREG 101 127 de 2026** (medidas ante El Niño, activada el 9 de septiembre de 2026 y vigente hasta "
      "mayo de 2027): despacho térmico mínimo obligatorio y una curva de referencia para las ofertas hidráulicas ligada "
      "al precio de escasez. Cambió la forma en que se forma el precio justo dentro del periodo de operación, y es la "
      "causa regulatoria más probable del aumento del error en septiembre y octubre (sección 9.5).")


def s3_datos():
    H("3. Datos e información disponible al corte")
    P("El corte es las **00:00 del día D**, y la ventana pronosticada va de la 01:00 de D a las 00:00 de D+1 (24 horas). "
      "Cada fuente entra al modelo solo con lo que ya estaba publicado a esa hora; esta regla se verificó midiendo la "
      "hora real de publicación en la API y en la biblioteca del portal de XM.")
    tabla(["Fuente", "Variables", "Resolución", "Disponible al corte"], [
        ["API de métricas de XM", "precio de bolsa, demanda, generación", "horaria", "hasta 23:00 de D-4 (~3 días de rezago)"],
        ["API de métricas de XM", "volumen útil de embalses, aportes hídricos", "diaria", "valor de D-2"],
        ["NOAA (CPC)", "índice ONI de El Niño / La Niña", "mensual (trimestre móvil)", "valor de M-2"],
        ["Portal de XM (iMAR)", "predespacho ideal: costo marginal por hora", "horaria", "el de D, publicado D-1 ~10:30"],
        ["Portal de XM (dMAR, dDEM)", "costo marginal y demanda del despacho programado", "horaria", "el de D, publicado D-1 ~14:30*"],
        ["Calendario", "hora, día de la semana, día del año, festivos de Colombia", "horaria", "conocido de antemano"],
    ], anchos=[2.6, 4.2, 2.1, 3.6], titulo="Tabla 3. Fuentes de datos y su disponibilidad",
        fuente="*Si a la hora de la corrida aún no está el despacho programado, se usa el del día anterior, igual que en el "
               "entrenamiento. Las ofertas de precio de los generadores (OFEI) no se usan: XM agrega los precios al mes "
               "siguiente, así que usarlas sería fuga de información.")
    figura(fig_calendario(), 16, "Figura 2. Qué información existe en el momento del pronóstico (corte a las 00:00 del día D).")
    P("Series del proyecto: un dataset maestro horario sincronizado de enero de 2019 a agosto de 2026 (66.576 horas) y "
      "una serie de operación que se actualiza todos los días con lo nuevo de XM y NOAA. El año 2026 se reservó como "
      "prueba fuera de muestra; 2019-2025 sirve para entrenar y para la validación histórica.")


def s4_variables():
    H("4. Variables de entrada (66)")
    P("Los modelos no predicen el precio directamente sino el **logaritmo de la razón entre el precio y el "
      "predespacho**, log(precio / predespacho). Las variables se agrupan en ocho familias. La notación de los conjuntos "
      "es la de los modelos: **R** = realistas (lo público al corte), **P** = predespacho, **F** = forma intradía del "
      "predespacho, **C** = despacho programado y **E** = error reciente del predespacho.")
    P("Un detalle importante en las variables de precio y demanda: como XM publica con ~3 días de rezago, lo que en el "
      "código se llama \"hace 24 h\" es en realidad **hace 96 h** (120 h para la hora 00:00, la última de la ventana). Se "
      "conservaron los nombres originales del proyecto para no romper el código, pero el valor ya está desplazado.")
    fam = [
        ("Tabla 4. Hidrología y clima (R) · 14 variables", [
            ["volumen_embalses", "Volumen útil de los embalses del SIN (valor diario de D-2)", "El agua guardada es el principal determinante del precio en un sistema hidro-dominado"],
            ["aportes_hidricos", "Caudal que llega a los embalses (D-2)", "Mide si está entrando agua; anticipa la oferta hidráulica"],
            ["…_delta_1d, …_delta_7d", "Cambio del embalse o de los aportes en 1 y 7 días", "Tendencia: un embalse que baja rápido encarece la energía"],
            ["…_media_7d, …_media_30d", "Promedio móvil de 7 y 30 días", "Nivel de referencia sin el ruido diario"],
            ["…_vs_media30d", "Valor actual / media de 30 días", "Si el sistema está más seco o más húmedo que en el último mes"],
            ["oni", "Índice Oceánico El Niño (NOAA), valor de hace dos meses", "El Niño reduce lluvias en Colombia y sube el precio con meses de anticipación"],
            ["es_pandemia", "1 en el periodo de restricciones COVID-19 (2020)", "Aísla la caída atípica de la demanda"],
        ]),
        ("Tabla 5. Precio pasado (R) · 9 variables", [
            ["precio_lag24h", "Precio de la misma hora hace 96 h (120 h para la 00:00)", "El precio más reciente publicado de esa hora"],
            ["precio_lag168h", "Precio de la misma hora una semana antes", "Patrón semanal (laboral contra fin de semana)"],
            ["precio_media_24h / 7d / 30d", "Promedios móviles del precio publicado", "Nivel reciente del mercado en tres escalas"],
            ["precio_std_24h / 7d", "Desviación estándar móvil", "Volatilidad reciente"],
            ["precio_rango_24h", "Máximo menos mínimo de 24 h", "Amplitud del último día publicado"],
            ["ratio_volatilidad", "std 24 h / std 7 d", "Si la volatilidad está subiendo o bajando"],
        ]),
        ("Tabla 6. Demanda (R) · 4 variables", [
            ["demanda_lag24h / 48h / 72h", "Demanda de la misma hora hace 96, 120 y 144 h", "Más demanda exige plantas más caras"],
            ["demanda_media_24h", "Promedio móvil de 24 h de la demanda publicada", "Nivel de consumo reciente"],
        ]),
        ("Tabla 7. Calendario (R) · 13 variables", [
            ["hora_sin, hora_cos", "Hora del día en forma circular (seno y coseno)", "La hora 23 queda cerca de la 0; el precio tiene picos de mañana y noche"],
            ["dia_semana_sin, dia_semana_cos", "Día de la semana circular", "Domingos y sábados tienen menos demanda"],
            ["dia_anio_sin, dia_anio_cos", "Día del año circular", "Estaciones secas y lluviosas"],
            ["es_festivo", "1 si el día es festivo en Colombia", "Un festivo se comporta como domingo"],
            ["festivo_lag24h / 48h / 72h / 168h", "Si fue festivo 1, 2, 3 o 7 días antes", "Efecto del festivo sobre los días siguientes"],
            ["mismatch_festivo_24h / 168h", "1 si hoy y el día de referencia (ayer o hace una semana) no coinciden en ser festivos", "Avisa que el precio de referencia no es comparable"],
        ]),
        ("Tabla 8. Predespacho ideal (P) · 9 variables", [
            ["pre_obj", "Costo marginal del predespacho en la hora pronosticada", "La mejor estimación pública del precio; es la referencia del modelo"],
            ["log_pre", "Logaritmo de pre_obj", "Escala en la que trabaja el modelo"],
            ["pre_media, pre_min, pre_max", "Promedio, mínimo y máximo del predespacho en el día", "Nivel y extremos que XM espera"],
            ["pre_rel", "pre_obj / pre_media", "Si la hora está por encima o por debajo del promedio del día"],
            ["pre_rango_rel", "(pre_max − pre_min) / pre_media", "Qué tan \"montañoso\" espera XM el día"],
            ["pre_lag24", "Predespacho de la misma hora el día anterior", "Cambio de un día a otro"],
            ["pre_cambio24", "pre_obj / pre_lag24", "Si XM espera un día más caro o más barato que ayer"],
        ]),
        ("Tabla 9. Forma intradía del predespacho (F) · 6 variables", [
            ["pre_d1", "(hora actual − hora anterior) / media del día", "Si la hora está en una rampa de subida o bajada"],
            ["pre_d1_sig", "(hora siguiente − hora actual) / media", "Hacia dónde va la rampa"],
            ["pre_rango_hora", "Posición de la hora dentro del día (0 = la más barata, 1 = la más cara)", "Ubica picos y valles"],
            ["pre_z", "(hora − media) / desviación del día", "Qué tan extrema es la hora dentro de su día"],
            ["pre_max_hora, pre_min_hora", "Hora en que el predespacho alcanza su máximo y su mínimo", "Momento esperado del pico y del valle"],
        ]),
        ("Tabla 10. Despacho programado (C) · 6 variables", [
            ["cm_programado", "Costo marginal del despacho programado (dMAR) en la hora", "Incluye restricciones de red que el predespacho ideal no tiene"],
            ["demanda_programada", "Demanda que XM programó para la hora (dDEM)", "Pronóstico oficial de demanda"],
            ["cm_prog_rel_pre", "cm_programado / pre_obj", "Diferencia entre los dos despachos de XM"],
            ["cm_prog_media", "Promedio diario del costo marginal programado", "Nivel del día según el programado"],
            ["dem_prog_rel, dem_prog_dia", "Demanda programada relativa a su promedio diario, y ese promedio", "Forma y nivel de la demanda esperada"],
        ]),
        ("Tabla 11. Error reciente del predespacho (E) · 5 variables", [
            ["err_dia_d4", "Error medio log(precio/predespacho) del último día completo publicado (D-5)", "Si XM viene sobreestimando o subestimando"],
            ["err_7d, err_30d", "Ese error promediado en 7 y 30 días", "Sesgo persistente del predespacho"],
            ["err_hora_d4, err_hora_7d", "El mismo error, solo para la hora pronosticada", "Horas en que el predespacho falla sistemáticamente"],
        ]),
    ]
    for tit, filas in fam:
        tabla(["Variable", "Definición", "Por qué se incluye"], filas, anchos=[3.3, 4.6, 4.6], tam=8.5, titulo=tit)
    nota("La familia P (predespacho) es la que más aporta: sin ella, con solo lo publicado al corte, el mejor ensamble del "
         "informe de avance apenas mejora a la persistencia (rMAE 0,93). El precio lo fija la oferta de la planta marginal, "
         "y el predespacho es el único dato público a tiempo que ya trae esas ofertas.", "Qué variables pesan más")


def s5_modelos():
    H("5. Modelos finales")
    H("5.1 Pronóstico de 24 horas", 2)
    P("**Paso 1. Tres CatBoost.** CatBoost es un modelo de árboles de decisión potenciados (gradient boosting) que maneja "
      "bien relaciones no lineales e interacciones sin necesidad de escalar las variables. Se entrenan tres, cada uno con un "
      "conjunto distinto de variables, para que se equivoquen en cosas diferentes:")
    tabla(["Modelo", "Variables", "Número"], [
        ["CatBoost R+P", "realistas + predespacho", "49"],
        ["CatBoost R+P+F+C", "+ forma intradía + despacho programado", "61"],
        ["CatBoost R+P+E", "+ error reciente del predespacho", "54"],
    ], anchos=[3, 6.5, 1.5], titulo="Tabla 12. Los tres CatBoost")
    P("Configuración común: 1.500 árboles, profundidad 6, tasa de aprendizaje 0,05, **pérdida MAE** (con lo que el modelo "
      "estima la mediana, que es robusta a los picos) y objetivo:")
    ec(r"y_t = \ln\left(\frac{\text{precio}_t}{\text{predespacho}_t}\right)",
       "La predicción se devuelve a pesos como predespacho × e^ŷ. Trabajar con la razón hace que el modelo aprenda "
       "correcciones relativas (por ejemplo, \"el precio real suele quedar 10 % por encima del predespacho en esta hora "
       "cuando el embalse está bajo\") que sirven igual con precios de 200 o de 800 COP/kWh.")
    P("Los tres se **reentrenan cada mes** con todo lo publicado hasta ese momento; en el entrenamiento solo entran "
      "objetivos cuyo precio tenía al menos 73 horas de publicado antes del corte. Luego se promedian.")
    P("**Paso 2. Combinador causal.** El promedio de los CatBoost y el predespacho se mezclan con una regresión "
      "cuantílica de la mediana (QRA, *quantile regression averaging*), una por cada franja de 6 horas del día:")
    ec(r"\hat{p}_t = w_0 + w_1\,\text{predespacho}_t + w_2\,\overline{\text{CatBoost}}_t",
       "Los pesos se reestiman cada día solo con días cuyo precio real ya estaba publicado: el combinador deja un margen "
       "de **5 días** entre lo que usa para aprender y el día que pronostica (al corte de D, la ventana de D-4 aún no está "
       "completa). En promedio el predespacho recibe un peso de 0,60 y los CatBoost de 0,37, de modo que cada pronóstico se "
       "puede leer como \"el predespacho de XM, corregido según el comportamiento reciente del mercado\".")
    P("**Paso 3. Bandas conformes del 80 %.** Para cada hora del día se toman los errores logarítmicos |ln(real/q50)| de "
      "esa misma hora en los 60 días anteriores con precio ya publicado (con el mismo margen de 5 días), y se calcula su "
      "cuantil del 80 % (con la corrección de muestra finita de la predicción conforme, nivel ⌈(n+1)·0,8⌉/n):")
    ec(r"s_i = \left| \ln \left( \text{real}_i / q50_i \right) \right|")
    ec(r"m = Q_{0{,}8}\left( s_i \right)")
    ec(r"q10 = q50 \cdot e^{-m}")
    ec(r"q90 = q50 \cdot e^{m}",
       "Es un método libre de supuestos de distribución: si en los últimos 60 días el pronóstico de las 19:00 se equivocó "
       "mucho, la banda de las 19:00 se abre; si acertó, se cierra.")
    H("5.2 Pronóstico de 72 horas", 2)
    P("Los pasos 1 a 24 son el pronóstico de 24 h. Para los pasos 25 a 72 (días D+1 y D+2) todavía no existe predespacho, "
      "así que la referencia es el predespacho del día D y se entrenan tres CatBoost por tramo de horizonte con la misma "
      "lógica. Las bandas se calibran igual, por horizonte.")
    H("5.3 Detector de picos y desplomes", 2)
    P("Un clasificador de árboles potenciados (HistGradientBoosting de scikit-learn, reentrenado cada semana) estima, para cada hora, la probabilidad de dos eventos definidos respecto a la mediana del "
      "precio de la ventana: **desplome** si el precio real queda en 0,7 veces la mediana o menos, y **pico** si queda en "
      "1,3 veces o más. Usa las mismas variables públicas al corte, incluido el predespacho. Sus probabilidades se muestran "
      "en el dashboard como alerta (una posibilidad, no una certeza); no modifican el pronóstico central.")
    P("Para el segundo y el tercer día del horizonte de 72 h hay un detector por día de anticipación, con lo publicado al "
      "corte (el predespacho del día del corte, no el del día pronosticado). Se validó en los cinco años históricos y en 2026 "
      "frente a la climatología (sección 9.7): el de **pico** se publica, marcado como menos confiable que el del primer día; "
      "el de **desplome** no, porque no supera a la climatología de forma consistente.")
    H("5.4 Modelos que se probaron y no quedaron", 2)
    P("El camino hasta el diseño final comparó persistencia, Prophet, XGBoost, ARX+GARCH, GARCH-ged (diseño LEAR), "
      "N-BEATSx y N-HiTS, y un ensamble de seis de ellos (el diseño del informe de avance). También se descartaron, por no "
      "mejorar o por usar información no disponible: XGBoost y LEAR sobre la razón con el predespacho, redes neuronales con "
      "predespacho, corrección de sesgo, ponderación por recencia, combinador por hora, las ofertas OFEI (fuga) y estirar "
      "la amplitud de los picos (DILATE). Las redes neuronales además resultaron muy sensibles a cambios pequeños en los "
      "datos. La sección 9.7 muestra la comparación completa, con validación en cinco años históricos, que justifica el "
      "diseño final.")


def s6_motor():
    H("6. Motor de decisión")
    P("El motor convierte el pronóstico en una señal por hora. Compara el precio esperado q50 con dos umbrales, por "
      "defecto los percentiles 25 (P25) y 75 (P75) del precio:")
    tabla(["Rol", "q50 ≥ P75 (precio alto)", "q50 ≤ P25 (precio bajo)", "En medio o poca confianza"], [
        ["Generador", "vender", "retener", "esperar"],
        ["Comercializador", "evitar la compra (usar contratos)", "comprar", "esperar"],
    ], anchos=[2.4, 3.5, 3.2, 3.2], titulo="Tabla 13. Señales del motor")
    tabla(["Método", "Cómo calcula los umbrales"], [
        ["fijo", "Percentiles del precio histórico 2019-2025, calculados una sola vez"],
        ["rodante", "Percentiles de los últimos 30 días, causales (solo pasado); se adaptan al régimen"],
        ["banda", "Como fijo, pero fuerza \"esperar\" cuando la banda q10-q90, relativa al precio, está por encima de su percentil 75 histórico (poca confianza)"],
        ["híbrido", "Umbral rodante + el filtro de confianza de banda"],
    ], anchos=[1.8, 10.2], titulo="Tabla 14. Los cuatro métodos de umbral")
    P("**Filtro intradía.** Los umbrales deciden sobre todo en qué temporada conviene la bolsa, pero casi no eligen la hora "
      "dentro del día. El filtro conserva la señal favorable (vender o comprar) solo en la mitad del día que el pronóstico "
      "pone más cara (generador) o más barata (comercializador), y la desfavorable en la mitad opuesta.")
    P("**Selección por estabilidad.** Cada método se evalúa por separado en las dos mitades del periodo. Es **estable** si "
      "en ambas mitades actúa entre el 10 % y el 40 % de las horas y su ventaja no es negativa. Entre los estables se elige "
      "el de mejor **peor mitad**, es decir, el que rinde mejor en su mitad más difícil (criterio conservador). Con los "
      "datos de 2026 el método elegido es el **rodante** para los dos roles y los dos horizontes.")


def s7_metricas():
    H("7. Métricas de calidad: qué son y por qué se usan")
    P("Ninguna métrica sola cuenta toda la historia. Augur usa un conjunto pequeño, cada una con una pregunta concreta. "
      "En las fórmulas, y es el precio real, ŷ el pronóstico y N el número de horas.")
    H("7.1 Error del pronóstico puntual", 2)
    P("**MAE, error absoluto medio (métrica principal).**")
    ec(r"\text{MAE} = \frac{1}{N}\sum_{t=1}^{N} \left| y_t - \hat{y}_t \right|",
       "Cuántos COP/kWh se equivoca el pronóstico en promedio. Se eligió como métrica principal porque está en las mismas "
       "unidades del precio (un comercializador lo entiende de inmediato), no se infla con unos pocos picos extremos y es "
       "coherente con el resto del sistema: los CatBoost y el combinador minimizan justamente el MAE, que corresponde a "
       "pronosticar la mediana. También es la métrica estándar en la literatura de pronóstico de precios de electricidad.")
    P("**RMSE, raíz del error cuadrático medio.**")
    ec(r"e_t = y_t - \hat{y}_t")
    ec(r"\text{MSE} = \frac{1}{N}\sum_{t=1}^{N} {e_t^{2}}")    # llaves: si no, el termino queda fuera de la sumatoria en OMML
    ec(r"\text{RMSE} = \sqrt{\text{MSE}}",
       "Penaliza más los errores grandes. Se reporta como complemento: si el RMSE es mucho mayor que el MAE, el error está "
       "concentrado en pocas horas (los picos).")
    P("**MAPE y sMAPE, errores porcentuales.**")
    ec(r"\text{MAPE} = \frac{100}{N}\sum_{t=1}^{N} \left| \left( y_t - \hat{y}_t \right) / y_t \right|")
    ec(r"\text{sMAPE} = \frac{100}{N}\sum_{t=1}^{N} \left( 2 \left| y_t - \hat{y}_t \right| / \left( \left| y_t \right| + \left| \hat{y}_t \right| \right) \right)",
       "Expresan el error como porcentaje del precio y permiten comparar con trabajos de otros periodos o países, donde el "
       "nivel de precios es distinto. El MAPE se dispara cuando el precio real es muy bajo; el sMAPE lo corrige "
       "dividiendo por el promedio de real y pronóstico.")
    P("**Sesgo.** Promedio de (pronóstico − real). Dice si el modelo tiende a quedarse corto o a pasarse; en 2026 el "
      f"sesgo del diseño definitivo es {n(K['sesgo'], 2)} COP/kWh, prácticamente cero.")
    H("7.2 Comparación contra una referencia", 2)
    P("**rMAE, MAE relativo a la persistencia.**")
    ec(r"\text{rMAE} = \frac{\text{MAE}_{\text{modelo}}}{\text{MAE}_{\text{persistencia}}}",
       "La persistencia es el pronóstico más simple posible: repetir el último precio publicado de la misma hora (en el "
       "escenario realista, el de hace 96 h). Un rMAE menor que 1 significa que el modelo aporta algo; 0,32 significa que "
       "comete el 32 % del error de la persistencia. Se usa porque un MAE de 30 COP/kWh es excelente en un mes tranquilo "
       "y mediocre en uno volátil: dividir por la persistencia del mismo periodo hace comparables meses, años y regímenes.")
    P("**Prueba de Diebold-Mariano (DM).** Un MAE más bajo puede ser suerte. La prueba DM compara, hora a hora, la "
      "diferencia de pérdidas entre dos pronósticos:")
    ec(r"d_t = \left| e^{A}_t \right| - \left| e^{B}_t \right|")
    ec(r"\text{DM} = \frac{\bar{d}}{s}",
       "donde d̄ es el promedio de las diferencias y s su error estándar. "
       "La varianza se estima con el método HAC (robusto a la autocorrelación), porque los errores de horas seguidas no "
       "son independientes. DM positivo significa que B se equivoca menos que A; con |DM| > 1,65 la diferencia es "
       "significativa al 5 % en una cola. Se exigió esta prueba para adoptar cualquier cambio: una variante solo entra si "
       "mejora con significancia.")
    H("7.3 Calidad de la incertidumbre", 2)
    P("**Cobertura.** Porcentaje de horas en que el precio real cayó dentro de la banda q10-q90. Como es una banda del "
      "80 %, lo ideal es que la cobertura sea cercana al 80 %: si es mucho menor, la banda promete más confianza de la que "
      f"tiene; si es mucho mayor, es innecesariamente ancha. En 2026: {pc(K['cobertura'])}.")
    P("**Ancho medio de la banda.** Promedio de q90 − q10. A igual cobertura, una banda más angosta es más útil.")
    P("**Puntaje de Winkler.** Une las dos anteriores en un solo número (menor es mejor):")
    ec(r"W_t = (q90_t - q10_t) + \frac{2}{\alpha}\left(q10_t - y_t\right)\mathbb{1}[y_t < q10_t] + \frac{2}{\alpha}\left(y_t - q90_t\right)\mathbb{1}[y_t > q90_t]",
       "con α = 0,2. Premia bandas angostas y castiga cada vez que el precio se sale, en proporción a cuánto se sale. Se usó "
       "para comparar variantes de bandas con la misma cobertura.")
    H("7.4 Calidad del detector de eventos", 2)
    V("**Puntaje de Brier:** error cuadrático medio de la probabilidad, (p − evento)², donde evento vale 0 o 1. Menor es mejor.")
    V("**BSS (Brier Skill Score):** mejora del Brier frente a la climatología (la frecuencia histórica del evento a esa hora). "
      "Positivo significa que el detector sabe más que \"lo que suele pasar\".")
    V("**AUC:** probabilidad de que el detector asigne más riesgo a una hora con evento que a una sin evento. 0,5 es azar, 1 es "
      "perfecto. Es independiente del umbral de alerta.")
    V("**Precisión promedio:** resume cuántas de las alertas son eventos reales; útil porque los eventos son raros "
      "(en 2026 hubo 566 horas de pico y 199 de desplome en 5.184).")
    H("7.5 Calidad de la decisión", 2)
    V("**Ventaja (COP/kWh):** para el generador, precio real promedio en las horas en que el motor dice vender menos el "
      "precio promedio del periodo; para el comercializador, al revés (promedio del periodo menos el precio en las horas de "
      "comprar). Positivo significa que la regla encuentra horas mejores que el promedio.")
    V("**Frecuencia de acción:** porcentaje de horas con señal de actuar. Se exige entre 10 % y 40 %: una regla que casi "
      "nunca actúa o que actúa siempre no sirve, aunque su ventaja se vea bien.")
    V("**Peor mitad y estabilidad:** la ventaja en la mitad más difícil del periodo; protege contra reglas que solo "
      "funcionaron en una racha.")
    nota("El objetivo del proyecto no es solo un MAE bajo: lo que le sirve al usuario es acertar las rampas y los cambios "
         "de nivel dentro del día. Por eso se miden también el error en las horas de rampa, la detección de eventos y la "
         "ventaja económica del motor. Un modelo con mejor MAE no siempre toma mejores decisiones.")


def s8_validacion():
    H("8. Cómo se validó")
    V("**Fuera de muestra:** todo 2026 (enero a agosto, 5.184 horas) se pronosticó sin haberlo usado para entrenar. "
      "La operación en vivo (desde el 5 de agosto) es una segunda prueba, completamente a ciegas.")
    V("**Causalidad estricta:** reentrenamiento mensual, combinador y bandas con margen de 5 días; en ningún punto se usa "
      "un precio que no estuviera publicado al corte. Se auditaron las fechas de publicación de cada fuente.")
    V("**Validación histórica (walk-forward):** cinco orígenes, uno por año de julio de 2020 a junio de 2025, que cubren "
      "La Niña, años neutros y el fuerte El Niño de 2023-2024. En cada uno se entrena solo con lo anterior y se pronostica "
      "el año siguiente. Sirve para confirmar que las decisiones tomadas mirando 2026 no son un sobreajuste.")
    V("**Prueba de significancia:** Diebold-Mariano con varianza HAC para cada comparación y para cada cambio adoptado.")
    V("**Referencias justas:** la persistencia y el predespacho se evalúan exactamente en las mismas horas y con la "
      "misma información que el modelo.")


def s9_resultados():
    H("9. Resultados")
    H("9.1 Pronóstico de 24 horas en 2026", 2)
    filas = []
    for f in T24:
        filas.append([f["modelo"], n(f["MAE"]), n(f["RMSE"]), pc(f["MAPE"]), n(f["rMAE"], 3),
                      "—" if f["DM_t"] is None or f["DM_t"] != f["DM_t"] else n(f["DM_t"], 2)])
    tabla(["Modelo", "MAE", "RMSE", "MAPE", "rMAE", "DM vs persistencia"], filas, anchos=[5.2, 1.3, 1.3, 1.3, 1.2, 1.9],
          titulo="Tabla 15. Todos los modelos con la misma información pública al corte (2026, 4.848 horas comunes)",
          fuente="La última fila es el diseño en producción. Sobre las 5.184 horas completas de 2026 su MAE es "
                 f"{n(K['MAE'])} COP/kWh y el MAPE {pc(K['MAPE'])}.")
    figura(fig_comparacion(), 15, "Figura 3. MAE de 24 h en 2026. Sin el predespacho, ningún modelo se aleja mucho de la "
                                  "persistencia; con él, el error cae a un tercio.")
    H("9.2 Pronóstico de 72 horas", 2)
    tabla(["Tramo", "MAE", "MAPE", "MAE persistencia", "rMAE", "Cobertura 80 %"], [
        [k, n(f["MAE"]), pc(f["MAPE"]), n(f["MAE_pers"]), n(f["rMAE"], 2), pc(f["cobertura"])] for k, f in T72.items()],
        anchos=[1.6, 1.3, 1.3, 2.2, 1.2, 2], titulo="Tabla 16. Error por tramo de horizonte (2026)",
        fuente="El error crece con el horizonte, como es de esperar, pero a 49-72 h el modelo sigue cometiendo el 61 % del "
               "error de la persistencia.")
    H("9.3 Por régimen climático", 2)
    rn, rx = C["regimen_normal_ene_abr"], C["regimen_el_nino_may_ago"]
    tabla(["Periodo", "Augur", "Predespacho", "Persistencia", "DM vs predespacho"], [
        ["Enero a abril (normal)", n(rn["MAE"]), n(rn["MAE_pre"]), n(rn["MAE_pers"]), n(rn["DM_vs_pre"][0], 2)],
        ["Mayo a agosto (El Niño)", n(rx["MAE"]), n(rx["MAE_pre"]), n(rx["MAE_pers"]), n(rx["DM_vs_pre"][0], 2)],
    ], anchos=[3.4, 1.6, 2, 2, 2.4], titulo="Tabla 17. MAE (COP/kWh) por régimen en 2026",
        fuente="En régimen normal Augur mejora con claridad al predespacho; con El Niño los dos quedan casi iguales "
               "(diferencia no significativa) y ambos superan ampliamente a la persistencia.")
    H("9.4 Validación histórica 2020-2025", 2)
    tabla(["Origen (año)", "Predespacho", "Versión simple", "Augur (v2)", "DM v2 vs simple"], [
        [f"{f['origen']}", n(f["MAE_predespacho"]), n(f["MAE_simple"]), n(f["MAE_v2"]), n(f["DM_t_v2_vs_simple"], 2)]
        for f in C["validacion_definitivo"]], anchos=[2.6, 2, 2.2, 2, 2.4],
        titulo="Tabla 18. MAE en cinco años históricos (jul-2020 a jun-2025)",
        fuente="Augur gana al predespacho en los cinco años. DM negativo = la versión v2 (tres CatBoost) se equivoca menos "
               "que la simple (un CatBoost); significativa en los orígenes 1 a 3.")
    H("9.5 Operación en vivo", 2)
    om = OP["mensual"]
    P(
      f"Desde el 5 de agosto el sistema corre solo todos los días. Entre esa fecha y el 11 de octubre (último precio "
      f"publicado al cierre de esta versión, {n(OP['horas'], 0)} horas) el MAE fue {n(OP['MAE'])} COP/kWh frente a "
      f"{n(OP['MAE_pre'])} del predespacho, y la cobertura de la banda {pc(OP['cobertura'])}. Por mes: agosto "
      f"{n(om['2026-08'])}, septiembre {n(om['2026-09'])} y octubre {n(om['2026-10'])}.")
    P("El salto de septiembre coincide con la activación de la Resolución CREG 101 127 (9 de septiembre): el error del "
      "predespacho de XM se duplicó en las mismas fechas, lo que indica un cambio en cómo se forma el precio y no una falla "
      "del modelo. Se probaron dos adaptaciones: dar más peso a los datos recientes (empeora) y agregar como variable una "
      "aproximación de la curva de referencia de las ofertas hidráulicas de la resolución. Esta última mejora en 2026, pero "
      "en los años históricos gana en 2 de 5, empata en 2 y pierde en 2024-25, así que no se adoptó.")
    H("9.6 Detector de eventos y motor", 2)
    ev = C["eventos"]
    tabla(["Evento", "Método", "Brier", "BSS", "AUC", "Detectados", "Falsas alarmas"], [
        [e["evento"], e["metodo"], n(e["Brier"], 3), n(e["BSS_vs_climatologia"], 2), n(e["AUC"], 2),
         f"{e['detectados_p05']} de {e['eventos']}", str(e["falsas_alarmas_p05"])] for e in ev],
        anchos=[1.5, 3.2, 1.1, 1, 1, 1.8, 1.6], tam=8.5, titulo="Tabla 19. Detección de picos y desplomes en 2026",
        fuente="Detectados y falsas alarmas con un umbral de probabilidad de 0,5.")
    m24 = C["motor"]
    tabla(["Rol y horizonte", "Método elegido", "Ventaja 1.ª mitad", "Ventaja 2.ª mitad"], [
        [f"{r} {h}", m24[f"{h}|{r}"]["metodo_estable"],
         n(m24[f"{h}|{r}"]["metodos"][m24[f"{h}|{r}"]["metodo_estable"]]["ventaja_h1"]),
         n(m24[f"{h}|{r}"]["metodos"][m24[f"{h}|{r}"]["metodo_estable"]]["ventaja_h2"])]
        for h in ("24h", "72h") for r in ("generador", "comercializador")],
        anchos=[3.4, 2.4, 2.4, 2.4], titulo="Tabla 20. Motor de decisión en 2026 (ventaja en COP/kWh)",
        fuente=f"Precio medio del periodo: {n(m24['precio_medio'])} COP/kWh. Las mitades se dividen el "
               f"{m24['corte_mitades']}. La ventaja no es una simulación de portafolio (sección 12).")
    s9_justificacion()


def s9_justificacion():
    """Benchmark de correctores, walk-forward de modelos desde cero y detector por horizonte (19 y 20-oct-2026)."""
    import pandas as pd
    base = RAIZ / "data/processed/resultados/escenario_realista"
    per = ["Origen 1 (2020-21)", "Origen 2 (2021-22)", "Origen 3 (2022-23)", "Origen 4 (2023-24)", "Origen 5 (2024-25)", "2026"]
    cab = ["2020-21", "2021-22", "2022-23", "2023-24", "2024-25", "2026"]
    H("9.7 Justificación del diseño: benchmark y walk-forward con información realista", 2)
    P("Cada componente se eligió con el mismo protocolo: comparar todos los candidatos razonables en las mismas horas, "
      "en cinco años históricos (15 de julio a 30 de junio, de 2020-21 a 2024-25) y en 2026, solo con la información "
      "pública al corte, y aceptar una diferencia solo si es significativa en la prueba de Diebold-Mariano.")
    b = pd.read_csv(base / "benchmark/metricas.csv")
    piv = b.pivot(index="modelo", columns="periodo", values="MAE")[per]
    orden = [("Persistencia", "Persistencia"), ("Predespacho", "Predespacho de XM solo"), ("Pre + CatBoost", "Predespacho + CatBoost (Augur)"),
             ("Pre + LightGBM", "Predespacho + LightGBM"), ("Pre + RandomForest", "Predespacho + Random Forest"),
             ("Pre + XGBoost", "Predespacho + XGBoost"), ("Pre + Lasso", "Predespacho + Lasso (lineal, tipo LEAR)")]
    tabla(["Modelo"] + cab + ["Media"], [[nom] + [n(piv.loc[k, c], 1) for c in per] + [n(piv.loc[k].mean(), 2)] for k, nom in orden],
          anchos=[4.2] + [1.1] * 6 + [1.1], tam=8.5,
          titulo="Tabla 21. MAE (COP/kWh) de los modelos que corrigen el predespacho",
          fuente="Mismas variables, objetivo, reentrenamiento mensual y presupuesto (1.500 árboles, profundidad 6), sin afinar "
                 "ningún modelo. Los cuatro de árboles no tienen diferencias significativas consistentes entre sí; el lineal es "
                 "significativamente peor en 3 de 6 periodos. CatBoost tiene la mejor media.")
    w = pd.read_csv(base / "walkforward/metricas.csv")
    pw = w.pivot(index="modelo", columns="periodo", values="MAE")[per]
    dm = w.pivot(index="modelo", columns="periodo", values="DM_t_vs_persistencia")[per]
    filas = []
    for k in ["Persistencia", "XGBoost", "ARX+GARCH", "GARCH-ged", "N-BEATSx", "N-HiTS", "Predespacho", "Pre + CatBoost"]:
        mejor = int(((dm.loc[k] < -1.96)).sum()) if k != "Persistencia" else None
        peor = int(((dm.loc[k] > 1.96)).sum()) if k != "Persistencia" else None
        nom = {"Predespacho": "Predespacho de XM solo", "Pre + CatBoost": "Predespacho + CatBoost (Augur)"}.get(k, k)
        filas.append([nom] + [n(pw.loc[k, c], 1) for c in per] + ["—" if mejor is None else f"{mejor} / {peor}"])
    tabla(["Modelo"] + cab + ["Mejor / peor"], filas, anchos=[3.6] + [1.05] * 6 + [1.5], tam=8.5,
          titulo="Tabla 22. Walk-forward con información realista: MAE (COP/kWh) y periodos en que cada modelo es "
                 "significativamente mejor o peor que la persistencia",
          fuente="2026 = 15 de enero a 5 de agosto. Sin el precio de ayer, ningún modelo que pronostica el precio desde cero "
                 "supera de forma consistente a la persistencia; el predespacho solo los supera a todos, y corregirlo lo mejora "
                 "en 5 de 6 periodos. En 2020-21, ARX+GARCH y las redes se entrenan con solo 17 meses y quedan fijos todo el "
                 "año, por eso su error es tan alto.")
    e = pd.read_csv(base / "eventos_72h/metricas.csv")
    e = e[e["metodo"] == "clasificador"]
    filas = []
    for ev in ("pico", "desplome"):
        for pasos, dia in (("1-24", "Día 1"), ("25-48", "Día 2"), ("49-72", "Día 3")):
            fila = [f"{ev.capitalize()} · {dia}"]
            for c in per:
                r = e[(e["evento"] == ev) & (e["pasos"] == pasos) & (e["periodo"] == c)].iloc[0]
                marca = "*" if (r["DM_p"] < 0.05 and r["DM_t_Brier_vs_clim"] < 0) else ("†" if r["DM_p"] < 0.05 else "")
                fila.append(f"{n(r['BSS_vs_climatologia'], 2)}{marca}")
            filas.append(fila)
    tabla(["Evento y día"] + cab, filas, anchos=[3.2] + [1.4] * 6, tam=8.5,
          titulo="Tabla 23. Detector de eventos por día del horizonte: mejora sobre la climatología (BSS)",
          fuente="BSS positivo = mejor que la frecuencia histórica de esa hora. * = significativamente mejor; † = "
                 "significativamente peor (Diebold-Mariano sobre el puntaje de Brier, p < 0,05). Se publican el día 1 "
                 "(pico y desplome) y el pico de los días 2 y 3.")


def s10_comparacion():
    H("10. Comparación con otros trabajos")
    P("La comparación numérica directa es difícil porque cada trabajo usa otra resolución (diaria, mensual u horaria), "
      "otro periodo y, sobre todo, otra información disponible. La tabla resume en qué se diferencia Augur.")
    tabla(["Trabajo", "Qué pronostica y cómo", "Error reportado", "Diferencia con Augur"], [
        ["Becerra Pulido (2025), USTA, Ing. Electrónica", "Precio diario 2024 con SARIMAX (demanda, hidráulica, embalses)",
         "MAPE 9,14 %", "Usa el precio del día anterior (no público a tiempo) y reemplaza los picos por la mediana; Augur es "
                        f"horario, sin esos supuestos: MAPE {pc(K['MAPE'])}"],
        ["Ferraro Rodríguez (2025), UNAL, maestría", "Correlaciones y ARIMA mensual 2010-2024", "no reporta métrica numérica",
         "Confirma que la oferta hidráulica explica el precio (r = 0,94); Augur lo aprovecha vía predespacho, que sí es público a tiempo"],
        ["García del Villar et al. (2020), Uninorte", "Precio promedio diario con Prophet, TBATS y redes",
         "la red subestimó 19 % en promedio", "Augur es horario, con bandas y motor de decisión"],
        ["Sánchez y Noguera (2024), Uninorte", "MPO ajustado (proyecto CREG 701 049) con redes neuronales", "dashboard diario",
         "Estudia la formación del precio; Augur pronostica el precio y recomienda acciones"],
        ["Barrientos et al. (2018), U. de Antioquia", "Precio horario 2000-2017 con redes, ARMA e híbridos", "MAPE 6,7 %",
         "Sin los eventos extremos de 2023-2026 y suponiendo conocido el precio del día anterior"],
        ["Kapoor y Wichitaksorn (2023), Nueva Zelanda", "33 modelos en un mercado hidro-dominado", "ninguno supera al ingenuo en todos los años",
         "Augur supera a su referencia ingenua en todos los años evaluados"],
        ["Predespacho ideal de XM", "Despacho del operador con la demanda pronosticada", f"MAE {n(K['MAE_pre'])} (2026)",
         f"Augur lo mejora en {pc(mejora_pre, 0)} en 2026 y en los cinco años históricos"],
    ], anchos=[3, 3.4, 2.2, 4.8], tam=8.5, titulo="Tabla 24. Antecedentes y comparación")
    nota("Varios antecedentes usan el precio de ayer como entrada, lo que hace sus errores más bajos en el papel pero "
         "irrealizables con datos públicos. Augur se evaluó con lo que de verdad existe a la hora de pronosticar.",
         "Lectura correcta")


def s11_operacion():
    H("11. Operación diaria")
    P("Un flujo de GitHub Actions corre tres veces al día (hora de Colombia, ~12:00, ~15:40 y ~18:40; se programan dos "
      "horas antes porque GitHub arranca con retraso). Cada corrida:")
    for i, t in enumerate(["Descarga lo nuevo de XM (precio, demanda, hidrología, predespacho, despacho programado) y el ONI de NOAA.",
                           "Si empezó un mes nuevo, reentrena los CatBoost (unos minutos en un servidor sin GPU).",
                           "Pronostica hasta la ventana del día siguiente (24 h) y los tres días siguientes (72 h), con bandas y probabilidades de eventos.",
                           "Recalcula el motor de decisión y sube los resultados al repositorio.",
                           "Streamlit Community Cloud detecta el cambio y redespliega el dashboard."], 1):
        p = doc.add_paragraph(style="List Number"); runs(p, t); p.paragraph_format.space_after = Pt(2)
    P("La corrida del mediodía ya tiene el predespacho del día siguiente; las de la tarde lo recalculan con el despacho "
      "programado. Solo se congelan los días que ya empezaron.")
    P("**Dashboard (augurpredictor.streamlit.app).** Tiene tres vistas: **Operador** (recomendación de la hora elegida, "
      "pronóstico del día con banda, horas para actuar y posibilidad de pico o desplome), **Analista** (semana, mapa de calor, "
      "backtest del motor y comparación de métodos) y **Metodología** (cómo funciona, fórmulas y validación). Las gráficas "
      "admiten zoom con Ctrl + rueda o el touchpad, desplazamiento, clic en una hora para analizarla y botón para centrar.")
    if (FIG_INF / "dash_operador.png").exists():
        figura(FIG_INF / "dash_operador.png", 15, "Figura 4. Vista Operador del dashboard.")


def s12_limitaciones():
    H("12. Limitaciones conocidas")
    for t in [
        "**Rezago de publicación.** El precio público llega con unos tres días de retraso; el modelo nunca ve el precio de ayer.",
        "**Dependencia del predespacho.** Los desfases de nivel de días completos que comete el predespacho (por ejemplo, del 11 al "
        "18 de junio de 2026) no se pueden anticipar con información pública, y Augur hereda la amplitud de los picos que el "
        "predespacho anticipa y no ocurren.",
        "**Cambios regulatorios.** La Resolución CREG 101 127 de 2026 cambió las reglas dentro del periodo de operación y duplicó "
        "el error de todos los pronósticos, incluido el del predespacho de XM. Las dos adaptaciones probadas (más peso a lo "
        "reciente y una aproximación de la curva de ofertas de la resolución) no pasaron la validación histórica.",
        "**PBN y PTB.** Augur pronostica el PBN; en horas de escasez el precio de liquidación (PTB) puede ser menor.",
        "**El motor no es una simulación de portafolio.** La ventaja compara precios promedio; no modela contratos firmados, "
        "capacidad de almacenamiento, costos de transacción ni la secuencia real de decisiones.",
        "**Rampas.** El modelo acierta la dirección de los cambios intradía mejor que su magnitud; el error se concentra en las "
        "horas de rampa (MAE ~117 COP/kWh en esas horas).",
        "**Picos y desplomes a 2 y 3 días.** Solo el detector de pico supera a la climatología en esos días, y no en todos los "
        "años; el de desplome no se publica.",
    ]:
        V(t)


def s13_glosario():
    H("13. Glosario")
    tabla(["Término", "Significado"], [
        ["Bolsa de energía", "Mercado de corto plazo donde se transa la energía no cubierta por contratos bilaterales"],
        ["PBN", "Precio de bolsa nacional: el precio horario que publica XM y que pronostica Augur"],
        ["PTB", "Precio de transacciones en bolsa, con el que se liquida; puede ser menor que el PBN en escasez"],
        ["MPO", "Máximo precio de oferta de las plantas despachadas en el despacho ideal"],
        ["Predespacho ideal", "Despacho que XM calcula el día anterior con la demanda pronosticada (archivo iMAR)"],
        ["Despacho programado", "Despacho con restricciones de red que XM publica D-1 ~14:30 (dMAR, dDEM)"],
        ["Corte", "Momento en que se hace el pronóstico: 00:00 del día pronosticado"],
        ["Persistencia", "Pronóstico ingenuo que repite el último precio publicado de la misma hora"],
        ["CatBoost", "Modelo de árboles de decisión potenciados (gradient boosting)"],
        ["QRA", "Combinación de pronósticos por regresión cuantílica"],
        ["Banda conforme", "Intervalo calibrado con los errores recientes, sin suponer una distribución"],
        ["q10, q50, q90", "Percentiles 10, 50 (mediana) y 90 del pronóstico"],
        ["ONI", "Índice Oceánico El Niño de la NOAA; positivo = El Niño, negativo = La Niña"],
        ["Walk-forward", "Validación que entrena con el pasado y prueba con el periodo siguiente, avanzando en el tiempo"],
        ["SIN", "Sistema Interconectado Nacional"],
        ["CREG", "Comisión de Regulación de Energía y Gas"],
    ], anchos=[3.2, 9.3], tam=9, titulo="Tabla 25. Términos usados en el documento")


SECCIONES = [s0_portada, s1_resumen, s2_mercado, s3_datos, s4_variables, s5_modelos, s6_motor, s7_metricas,
             s8_validacion, s9_resultados, s10_comparacion, s11_operacion, s12_limitaciones, s13_glosario]
for i, f in enumerate(SECCIONES):
    if i <= HASTA:
        f()

# pie de pagina con numero
pie = sec.footer.paragraphs[0]
pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = pie.add_run("Augur · funcionamiento del sistema · página ")
r.font.size, r.font.color.rgb = Pt(8.5), GRIS
for tipo, txt in (("begin", None), (None, "PAGE"), ("end", None)):
    rr = pie.add_run(); rr.font.size = Pt(8.5)
    if tipo:
        fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tipo); rr._r.append(fc)
    else:
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; rr._r.append(it)

doc.save(str(DESTINO))
print("ok", DESTINO, "secciones hasta", min(HASTA, len(SECCIONES) - 1))
