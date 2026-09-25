# -*- coding: utf-8 -*-
"""Actualiza SOLO la parte del motor de decision (OE3/OE4) del informe de avance ABET.

Copia docs/Informe_Avance_ABET_Barcelo_Dede.docx a una copia LOCAL fuera del
repositorio (la carpeta que contiene al repo) y edita esa copia: cambia los parrafos,
filas de tabla y referencias que describen el motor e inserta las tablas y figuras
nuevas copiando el formato de las existentes. El original del repo no se toca.
Los parrafos se ubican por su texto, no por posicion. Los numeros salen de
informe_oe3_assets/resultados_oe3.json (generado por informe_avance_oe3.py).
"""
import copy
import json
import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.table import Table
from docx.text.paragraph import Paragraph

RAIZ = Path(__file__).resolve().parents[1]
ASSETS = Path(__file__).resolve().parent / "informe_oe3_assets"
ORIGEN = RAIZ / "docs" / "Informe_Avance_ABET_Barcelo_Dede.docx"
DESTINO = RAIZ.parent / "Informe_Avance_ABET_Barcelo_Dede_OE3.docx"

R = json.loads((ASSETS / "resultados_oe3.json").read_text(encoding="utf-8"))
shutil.copyfile(ORIGEN, DESTINO)
doc = Document(str(DESTINO))
body = doc.element.body


def f1(v):
    return f"{v:,.1f}".replace(",", "X").replace(".", ",").replace("X", ".").replace("-", "−")


def pct(v):
    return f"{v * 100:.1f}".replace(".", ",") + " %"


def buscar_parrafo(prefijo):
    for el in body.iterchildren():
        if el.tag == qn("w:p") and Paragraph(el, doc).text.strip().startswith(prefijo):
            return el
    raise LookupError(f"No encontre el parrafo que empieza por: {prefijo!r}")


def siguiente_tabla(el):
    nxt = el.getnext()
    while nxt is not None and nxt.tag != qn("w:tbl"):
        nxt = nxt.getnext()
    return nxt


def reescribir(p_el, segmentos):
    """Reemplaza el texto del parrafo por segmentos (texto, negrita), conservando su
    formato de parrafo y el formato de caracter de su primer run."""
    rpr_base = None
    primer_r = p_el.find(qn("w:r"))
    if primer_r is not None and primer_r.find(qn("w:rPr")) is not None:
        rpr_base = copy.deepcopy(primer_r.find(qn("w:rPr")))
        for tag in ("w:b", "w:bCs", "w:i", "w:iCs"):
            for x in rpr_base.findall(qn(tag)):
                rpr_base.remove(x)
    for child in list(p_el):
        if child.tag in (qn("w:r"), qn("w:proofErr")):
            p_el.remove(child)
    par = Paragraph(p_el, doc._body)
    for seg in segmentos:
        texto, negrita = seg[0], seg[1]
        cursiva = seg[2] if len(seg) > 2 else False
        run = par.add_run(texto)
        if rpr_base is not None:
            run._r.insert(0, copy.deepcopy(rpr_base))
        if negrita:
            run.bold = True
        if cursiva:
            run.italic = True
    return p_el


def texto_celda(tc, texto):
    p_el = tc.find(qn("w:p"))
    runs = p_el.findall(qn("w:r"))
    for r in runs[1:]:
        p_el.remove(r)
    if runs:
        for t in runs[0].findall(qn("w:t")):
            runs[0].remove(t)
        t = runs[0].makeelement(qn("w:t"), {})
        t.text = texto
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        runs[0].append(t)
    else:
        Paragraph(p_el, doc._body).add_run(texto)


def insertar_despues(ancla, nuevo):
    ancla.addnext(nuevo)
    return nuevo


def parrafo_imagen(ancla, plantilla_img, archivo, ancho_in):
    p_el = copy.deepcopy(plantilla_img)
    for child in list(p_el):
        if child.tag != qn("w:pPr"):
            p_el.remove(child)
    insertar_despues(ancla, p_el)
    Paragraph(p_el, doc._body).add_run().add_picture(str(ASSETS / archivo), width=Inches(ancho_in))
    return p_el


def pie(ancla, plantilla_pie, etiqueta, texto):
    return insertar_despues(ancla, reescribir(copy.deepcopy(plantilla_pie), [(etiqueta, True), (" " + texto, False)]))


def parrafo_nuevo(ancla, plantilla, segmentos):
    return insertar_despues(ancla, reescribir(copy.deepcopy(plantilla), segmentos))


casos = R["casos"]
g24, c24 = casos["24h|generador"], casos["24h|comercializador"]
g72, c72 = casos["72h|generador"], casos["72h|comercializador"]
um = R["ultimo_mes"]["metodos"]
bar = {r["par"]: r for r in R["barrido_percentiles"]}
corte_txt = "19 de abril" if R["corte"] == "2026-04-19" else R["corte"]

# ---------------------------------------------------------------------------
# 6.8 Diseno definitivo: descripcion del motor + parrafo nuevo del dashboard
# ---------------------------------------------------------------------------
p_motor = buscar_parrafo("Motor de decisión (método «banda»).")
reescribir(p_motor, [
    ("Motor de decisión (cuatro métodos y selección por estabilidad).", True),
    (" (1) Cada método fija un umbral bajo y uno alto y compara con ellos la mediana q50 de cada hora: "
     "por debajo del bajo el comercializador recibe «comprar» y el generador «retener»; por encima del "
     "alto, «evitar compra» y «vender»; en otro caso, «esperar». Los umbrales por defecto son los "
     "percentiles 25 y 75, el rango intercuartílico convencional para separar valores típicos de "
     "extremos [34]. (2) Los métodos difieren en la referencia: «fijo» usa los percentiles del precio de "
     "entrenamiento 2019-2025; «rodante», los de una ventana móvil causal de 30 días sobre q50, que se "
     "adapta al régimen vigente [35]; «banda» es el fijo más un filtro que emite «esperar» cuando el ancho "
     "q90 − q10 supera su percentil 75 histórico; e «híbrido», agregado en esta etapa, combina el umbral "
     "rodante con ese filtro de confianza. (3) El método activo se elige por backtesting: solo son "
     "admisibles los que actúan entre 10 % y 40 % de las horas y, entre ellos, gana el de mayor ventaja "
     "en la mitad más débil del periodo (criterio maximin [36]) y no el de mayor ventaja promedio, para "
     "descartar un método cuyo buen promedio dependa de un solo tramo del año.", False),
])
parrafo_nuevo(p_motor, p_motor, [
    ("Dashboard.", True),
    (" La vista Operador aplica el método que elige el criterio de estabilidad y muestra, para la hora "
     "seleccionada, la acción recomendada, la posición del precio esperado entre los umbrales y un nivel "
     "de confianza según el ancho de la banda; la hora se elige con un selector o haciendo clic sobre la "
     "franja de 24 horas del día. La vista Analista permite fijar a mano el método (incluido «híbrido») y "
     "los percentiles, y recalcula el método sugerido sobre el rango de fechas elegido. Ambas vistas se "
     "verificaron sin errores con pruebas automatizadas de Streamlit en las cuatro combinaciones de rol "
     "y horizonte.", False),
])

# Tabla 13 (decisiones de diseno) y Tabla 10 (riesgos): solo las filas del motor
for tbl in body.iter(qn("w:tbl")):
    for tr in tbl.iter(qn("w:tr")):
        tcs = tr.findall(qn("w:tc"))
        textos = ["".join(t.text or "" for t in tc.iter(qn("w:t"))) for tc in tcs]
        if textos and textos[0].startswith("Motor elegido por backtest"):
            texto_celda(tcs[0], "Motor elegido por backtest: frecuencia de acción del 10 al 40 % y criterio "
                                "de estabilidad (maximin entre mitades del periodo)")
            texto_celda(tcs[1], "Elegir por la ventaja promedio del año; un umbral fijo único")
            texto_celda(tcs[2], f"«banda» gana el promedio pero actúa en el {pct(g24['metodos']['banda']['frec_h1'])} "
                                "de las horas de la primera mitad; el criterio estable elige «híbrido» "
                                "(generador) y «rodante» (comercializador) (Tablas 18 y 19)")
        if textos and textos[-1].startswith("Método «banda»: el motor espera"):
            texto_celda(tcs[-1], "Métodos «banda» e «híbrido»: el motor espera cuando la incertidumbre es alta.")

# ---------------------------------------------------------------------------
# 7.3 Resultados: parrafo del motor, Tabla 18 actualizada, Tabla 19 y Figuras 5-7
# ---------------------------------------------------------------------------
p_res = buscar_parrafo("OE3 y OE4 — Motor de decisión.")
reescribir(p_res, [
    ("OE3 y OE4 — Motor de decisión.", True),
    (" La ventaja es la diferencia entre el precio medio en las horas en que el motor recomienda actuar "
     "y el del periodo (precio medio 2026: 397,9 COP/kWh). Con el criterio de ventaja promedio, «banda» "
     f"gana en los cuatro casos ({f1(g24['metodos']['banda']['ventaja_anio'])}, "
     f"{f1(c24['metodos']['banda']['ventaja_anio'])}, {f1(g72['metodos']['banda']['ventaja_anio'])} y "
     f"{f1(c72['metodos']['banda']['ventaja_anio'])} COP/kWh). Al evaluar por separado las dos mitades del "
     f"periodo (corte: {corte_txt}) ese resultado no se sostiene: para el generador a 24 h, «banda» actúa en "
     f"el {pct(g24['metodos']['banda']['frec_h1'])} de las horas de la primera mitad y en el "
     f"{pct(g24['metodos']['banda']['frec_h2'])} de la segunda, y para el comercializador no actúa en "
     "ninguna hora de la segunda mitad (Tabla 19, Figura 5). La causa es la referencia fija: el precio de "
     "2026 se alejó del rango 2019-2025, de modo que casi todas las horas quedan por encima o por debajo "
     "de los percentiles históricos. Con el criterio de estabilidad el motor elige «híbrido» para el "
     "generador y «rodante» para el comercializador en ambos horizontes, los únicos métodos con ventaja "
     "positiva y frecuencia válida en las dos mitades (Tabla 18). Su ventaja promedio es menor "
     f"({f1(g72['metodos']['hibrido']['ventaja_anio'])} a {f1(c24['metodos']['rodante']['ventaja_anio'])} "
     "COP/kWh), pero se mantiene en todo el año.", False),
])

cap18 = buscar_parrafo("Tabla 18.")
reescribir(cap18, [("Tabla 18.", True),
                   (" Backtesting del motor de decisión sobre 2026: método elegido por el criterio de estabilidad.", False)])
tbl18 = siguiente_tabla(cap18)
filas18 = tbl18.findall(qn("w:tr"))
enc = filas18[0].findall(qn("w:tc"))
texto_celda(enc[3], "Ventaja 1.ª / 2.ª mitad (COP/kWh)")
texto_celda(enc[4], "Frecuencia 1.ª / 2.ª mitad")
for tr, (h, rol, c) in zip(filas18[1:], [("24h", "generador", g24), ("24h", "comercializador", c24),
                                          ("72h", "generador", g72), ("72h", "comercializador", c72)]):
    m = c["metodo_estable"]
    d = c["metodos"][m]
    celdas = tr.findall(qn("w:tc"))
    texto_celda(celdas[2], "híbrido" if m == "hibrido" else m)
    texto_celda(celdas[3], f"{f1(d['ventaja_h1'])} / {f1(d['ventaja_h2'])}")
    texto_celda(celdas[4], f"{pct(d['frec_h1'])} / {pct(d['frec_h2'])}")

# Tabla 19: los 4 metodos a 24 h, promedio del anio vs. mitades (7 columnas)
plantilla_pie_tabla = cap18
cap19 = pie(tbl18, plantilla_pie_tabla, "Tabla 19.",
            "Métodos del motor a 24 h: ventaja del año completo frente a cada mitad del periodo.")
tbl19 = copy.deepcopy(tbl18)
grid = tbl19.find(qn("w:tblGrid"))
for gc in list(grid):
    grid.remove(gc)
anchos = [1500, 1150, 1250, 1100, 1650, 1600, 1100]
for w in anchos:
    gc = grid.makeelement(qn("w:gridCol"), {qn("w:w"): str(w)})
    grid.append(gc)
trs = tbl19.findall(qn("w:tr"))
tr_enc, tr_dato = trs[0], trs[1]
for tr in trs[1:]:
    tbl19.remove(tr)
for tr in (tr_enc, tr_dato):
    tcs = tr.findall(qn("w:tc"))
    for _ in range(len(anchos) - len(tcs)):
        tr.append(copy.deepcopy(tcs[-1]))
    for tc, w in zip(tr.findall(qn("w:tc")), anchos):
        tc.find(qn("w:tcPr")).find(qn("w:tcW")).set(qn("w:w"), str(w))
for tc, t in zip(tr_enc.findall(qn("w:tc")), ["Rol", "Método", "Ventaja año (COP/kWh)", "Frecuencia año",
                                               "Ventaja 1.ª / 2.ª mitad", "Frecuencia 1.ª / 2.ª mitad", "¿Estable?"]):
    texto_celda(tc, t)


def vtxt(v):
    return "sin acción" if v is None else f1(v)


for rol, c in [("generador", g24), ("comercializador", c24)]:
    for m in ["fijo", "rodante", "banda", "hibrido"]:
        d = c["metodos"][m]
        tr = copy.deepcopy(tr_dato)
        valores = [rol, "híbrido" if m == "hibrido" else m, vtxt(d["ventaja_anio"]), pct(d["frec_anio"]),
                   f"{vtxt(d['ventaja_h1'])} / {vtxt(d['ventaja_h2'])}",
                   f"{pct(d['frec_h1'])} / {pct(d['frec_h2'])}", "sí" if d["estable"] else "no"]
        for tc, v in zip(tr.findall(qn("w:tc")), valores):
            texto_celda(tc, v)
        tbl19.append(tr)
insertar_despues(cap19, tbl19)

# Figuras 5, 6 y 7 (formato copiado de la Figura 4)
cap4 = buscar_parrafo("Figura 4.")
img4 = cap4.getprevious()
p_img5 = parrafo_imagen(tbl19, img4, "oe3_frecuencia_mitades.png", 4.9)
cap5 = pie(p_img5, cap4, "Figura 5.",
           "Horas con acción del generador a 24 h en cada mitad del periodo 2026, por método. "
           "Franja verde: rango válido de 10 a 40 %.")

p_sens = parrafo_nuevo(cap5, p_res, [
    ("Sensibilidad de los umbrales.", True),
    (" Se barrieron seis pares de percentiles (10/90 a 35/65) con el método fijo para el generador a 24 h "
     "(Figura 6). No aparece un óptimo interior: los pares más extremos dan más ventaja por hora de acción "
     f"pero actúan menos ({f1(bar['10/90']['ventaja'])} COP/kWh en el {f1(bar['10/90']['frec'])} % de las horas "
     f"con 10/90, frente a {f1(bar['35/65']['ventaja'])} COP/kWh en el {f1(bar['35/65']['frec'])} % con 35/65). "
     f"El par 25/75 queda en un punto intermedio ({f1(bar['25/75']['ventaja'])} COP/kWh, "
     f"{f1(bar['25/75']['frec'])} % de las horas), y para el comercializador la ventaja casi no cambia con el "
     "par elegido (272,7 a 286,4 COP/kWh). En el barrido documentado en la bitácora también se probaron "
     "ventanas rodantes de 7 a 90 días y percentiles del filtro de ancho de 60 a 90: la ventana de 30 días "
     "fue más pareja entre mitades que la de 90, y el filtro mostró una curva plana alrededor de 75, por lo "
     "que se conservaron los valores por defecto.", False),
])
p_img6 = parrafo_imagen(p_sens, img4, "oe3_barrido_percentiles.png", 4.9)
cap6 = pie(p_img6, cap4, "Figura 6.",
           "Ventaja frente a frecuencia de acción para seis pares de percentiles (método fijo, generador, 24 h).")

mes = R["ultimo_mes"]
p_mes = parrafo_nuevo(cap6, p_res, [
    ("Último mes y traducción a pesos.", True),
    (f" Entre el 7 de julio y el 5 de agosto de 2026 la banda del pronóstico cubrió el {f1(mes['cobertura'])} % "
     f"de las horas (objetivo 80 %) y el MAE subió a {f1(mes['mae'])} COP/kWh. Para el generador a 24 h, "
     f"«banda» dio una ventaja de {f1(um['banda']['ventaja'])} COP/kWh, «fijo» actuó en el 100 % de las horas "
     f"(ventaja nula), «híbrido» dio +{f1(um['hibrido']['ventaja'])} COP/kWh en el {um['hibrido']['frec']:.0f} % de "
     f"las horas y «rodante» +{f1(um['rodante']['ventaja'])} COP/kWh en el {um['rodante']['frec']:.0f} %. Para un "
     f"cliente de ejemplo de {R['potencia_kw']} kW que opere a potencia constante en las horas de acción, esto "
     f"equivale a {um['banda']['cop'] / 1e6:+.2f} millones de COP con «banda», "
     f"{um['hibrido']['cop'] / 1e6:+.2f} millones con «híbrido» y {um['rodante']['cop'] / 1e6:+.2f} millones con "
     "«rodante» (Figura 7). Al usar ese mismo mes como referencia en lugar de 2019-2025, el método fijo "
     "también da ventaja positiva con el mismo pronóstico: la pérdida de «banda» se debe sobre todo a la "
     "referencia histórica fija y no al modelo de pronóstico. La traducción a pesos es ilustrativa y no "
     "sustituye una evaluación económica completa [15].".replace("+0.", "+0,").replace("-0.", "−0,")
     .replace("+1.", "+1,").replace("+2.", "+2,"), False),
])
p_img7 = parrafo_imagen(p_mes, img4, "oe3_ultimo_mes_pesos.png", 4.9)
pie(p_img7, cap4, "Figura 7.",
    f"Ganancia o pérdida total por método en el último mes (generador, 24 h, cliente de ejemplo de "
    f"{R['potencia_kw']} kW).")

# ---------------------------------------------------------------------------
# 7.4 Evidencia por objetivo, 8.4 Impacto economico, 9 Conclusiones
# ---------------------------------------------------------------------------
p_ev = buscar_parrafo("Evidencia por objetivo.")
texto_ev = Paragraph(p_ev, doc).text
viejo = "OE3: motor con tres métodos, dos roles y dos horizontes (Tabla 18, Figura 3). OE4: backtesting hecho; falta la validación con usuarios."
nuevo = ("OE3: motor con cuatro métodos, selección por estabilidad, dos roles y dos horizontes (Tablas 18 y 19, "
         "Figuras 3 y 5 a 7). OE4: backtesting del año completo y por mitades hecho; falta la validación con usuarios.")
if viejo not in texto_ev:
    raise LookupError("No encontre la frase de OE3/OE4 en 'Evidencia por objetivo'")
resto = texto_ev.replace("Evidencia por objetivo.", "", 1).replace(viejo, nuevo)
reescribir(p_ev, [("Evidencia por objetivo.", True), (resto, False)])

p_eco = buscar_parrafo("El costo del proyecto es de")
texto_eco = Paragraph(p_eco, doc).text
primera = texto_eco.split(" En el backtesting")[0]
reescribir(p_eco, [(
    primera + " En el backtesting de 2026 a 24 h, con el método que elige el criterio de estabilidad, las horas en "
    f"que el motor recomienda vender tuvieron un precio medio {f1(g24['metodos']['hibrido']['ventaja_anio'])} COP/kWh "
    f"por encima del promedio, y las de comprar, {f1(c24['metodos']['rodante']['ventaja_anio'])} COP/kWh por debajo. "
    f"Para un cliente de ejemplo de {R['potencia_kw']} kW, en el último mes evaluado esto equivale a "
    f"{um['hibrido']['cop'] / 1e6:+.2f} millones de COP con «híbrido», frente a {um['banda']['cop'] / 1e6:+.2f} "
    "millones con «banda». No es una ganancia garantizada: no modela contratos, capacidad ni costos de "
    "transacción.".replace("+1.", "+1,").replace("-0.", "−0,"), False)])

reescribir(buscar_parrafo("OE3 (70 %)"), [
    ("OE3 (70 %): en ejecución.", True),
    (" El motor traduce las bandas en señales de comprar, vender o esperar para dos roles y dos horizontes "
     "con cuatro métodos, y elige el método activo por un criterio de estabilidad entre las dos mitades del "
     "periodo; el dashboard funciona y se verificó con pruebas automatizadas. Faltan la biblioteca de "
     "imágenes, conectar el motor al ensamble de seis votantes (hoy usa las bandas de N-BEATSx) e integrar "
     "las métricas de forma del día [15].", False),
])
reescribir(buscar_parrafo("OE4 (35 %)"), [
    ("OE4 (35 %): en ejecución.", True),
    (" El backtesting por mitades mostró que el método de mejor promedio («banda») no se sostiene en todo "
     "2026; el criterio de estabilidad elige «híbrido» y «rodante», con ventaja positiva y frecuencia válida "
     "en ambas mitades. Falta la validación con 3 a 5 usuarios, planeada en la Fase 4.", False),
])

rec5 = buscar_parrafo("5.\tEn una fase posterior")
parrafo_nuevo(rec5, rec5, [
    ("6.", False),
    ("\tEvaluar el criterio de estabilidad con más de dos subperiodos o con una formulación de bandits no "
     "estacionarios [37], y probar formalmente la ventana de 30 días y los percentiles 25/75 del motor.", False),
])

# ---------------------------------------------------------------------------
# Bibliografia: [34] a [37], mismo formato IEEE que las anteriores
# ---------------------------------------------------------------------------
ancla = buscar_parrafo("[33]")
nuevas = [
    [("[34]", False), ("\tJ. W. Tukey, ", False), ("Exploratory Data Analysis", False, True),
     (". Reading, MA, EE. UU.: Addison-Wesley, 1977.", False)],
    [("[35]", False), ("\tJ. Gama, I. Žliobaitė, A. Bifet, M. Pechenizkiy y A. Bouchachia, \"A survey on concept "
                       "drift adaptation,\" ", False), ("ACM Comput. Surv.", False, True),
     (", vol. 46, no. 4, art. 44, 2014, doi: 10.1145/2523813.", False)],
    [("[36]", False), ("\tA. Wald, ", False), ("Statistical Decision Functions", False, True),
     (". Nueva York, NY, EE. UU.: Wiley, 1950.", False)],
    [("[37]", False), ("\tA. Garivier y E. Moulines, \"On upper-confidence bound policies for switching bandit "
                       "problems,\" en ", False), ("Algorithmic Learning Theory (ALT 2011)", False, True),
     (", LNCS, vol. 6925. Berlín, Alemania: Springer, 2011, pp. 174–188.", False)],
]
for segs in nuevas:
    ancla = parrafo_nuevo(ancla, ancla, segs)

doc.save(str(DESTINO))
print("Guardado:", DESTINO)
