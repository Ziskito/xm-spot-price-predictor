# -*- coding: utf-8 -*-
"""
Genera la tabla de riesgos del proyecto en Excel, mismo formato que los dos ejemplos del
profesor (encabezado azul, texto blanco en negrita, columnas: Tipo de Riesgo, Descripcion,
Impacto, Descripcion del Impacto, Probabilidad, Medida de mitigacion).

Riesgos elegidos: varios tomados directamente de la seccion 6.4 del Anexo 1 (factores SO2),
otros de cosas que realmente ya ocurrieron o se identificaron en el desarrollo del proyecto
(fugas de datos, sobrecalentamiento del equipo, hallazgos empiricos sobre El Nino).
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pathlib import Path

RAIZ = Path(r"C:\Users\mgdbj\xm-spot-price-predictor")

FILAS = [
    ("Administrativo",
     "Incumplimiento de las fechas del cronograma, en particular Fase 3 (comparacion de modelos + motor de decision), que vence el 20 de septiembre y coincide con volumen alto de trabajo pendiente en OE3/OE4.",
     "Alto",
     "Retraso en el informe comparativo y en el arranque del motor de decision, con efecto en cascada sobre OE4 (validacion) y las entregas de la semana 16.",
     "Media",
     "Control semanal de avance contra el cronograma; priorizar el cierre de OE3 sobre exploracion adicional de modelos una vez la comparacion formal este consolidada."),

    ("Tecnico / Regulatorio",
     "Intervencion regulatoria transitoria de la CREG en la formacion del precio de bolsa (p. ej. techos a ofertas en periodos de escasez), reconocida explicitamente como riesgo en el Anexo 1 (6.3.4).",
     "Alto",
     "Un cambio en las reglas de formacion del precio a mitad de proyecto puede volver obsoletos los patrones aprendidos por los modelos entrenados con el historico anterior.",
     "Media",
     "Documentar la fecha de corte de los datos de entrenamiento en cada entrega; monitorear anuncios oficiales de la CREG; disenar el pipeline para poder reentrenar rapido si cambia el regimen."),

    ("Tecnico",
     "Cambios o restriccion de acceso a las APIs publicas de XM/SIMEM sin previo aviso (riesgo reconocido en el Anexo 1, 6.4.2).",
     "Alto",
     "Bloqueo de la actualizacion del dataset y del pipeline automatizado de adquisicion de datos.",
     "Baja",
     "Mantener copias locales versionadas de los datasets ya descargados (en el repositorio de GitHub); aislar el codigo de descarga en modulos independientes, faciles de adaptar si cambia el formato del API."),

    ("Tecnico",
     "Fuga de informacion (data leakage): variables que no respetan el horizonte de pronostico de 24-72h, dando al modelo acceso a datos que no existirian en un pronostico real.",
     "Alto",
     "Invalida silenciosamente las metricas de precision reportadas -- ya ocurrio mas de una vez durante el desarrollo (ej. hidrologia y precio rezagado en Prophet) y se corrigio al detectarse.",
     "Media",
     "Funcion central de verificacion de rezagos por variable antes de entrenar; revision cruzada entre los dos integrantes de cualquier feature nueva antes de reportar resultados finales."),

    ("Tecnico",
     "Los modelos evaluados (incluyendo arquitecturas de deep learning) no logran mejorar de forma estadisticamente significativa el error de pronostico durante el regimen El Nino, el escenario de mayor riesgo economico para el usuario.",
     "Medio",
     "No invalida el proyecto, pero limita el valor practico del motor de decision justo en el escenario donde mas se necesita precision.",
     "Alta",
     "Disenar el motor de decision (OE3) para ser explicitamente mas conservador (bandas de incertidumbre mas anchas, umbrales de accion mas estrictos) durante El Nino, en vez de depender de que el modelo de precio mejore su precision puntual ahi."),

    ("Operativo",
     "Sobrecalentamiento del equipo de computo durante entrenamientos prolongados de los modelos de deep learning (temperaturas de hasta 96 grados observadas en sesiones de entrenamiento sostenido).",
     "Alto",
     "Interrupcion del entrenamiento en curso y riesgo de dano al hardware si se sostiene por periodos largos.",
     "Media",
     "Secuenciar los entrenamientos pesados (uno a la vez, nunca en paralelo); monitorear temperatura durante sesiones largas; evaluar migrar cargas pesadas a un equipo con GPU dedicada si persiste."),

    ("Operativo",
     "Perdida o corrupcion de datasets procesados o resultados de experimentos de modelado, algunos de los cuales toman horas de computo en regenerarse.",
     "Alto",
     "Perdida de tiempo de computo significativo y de resultados de experimentos ya validados.",
     "Baja",
     "Control de versiones en GitHub con commits frecuentes de datasets y resultados intermedios (ya vigente); estructura de carpetas separada para datos crudos, procesados y resultados."),

    ("Social / Equipo",
     "Descoordinacion en el reparto de trabajo entre los dos integrantes (modelado en OE2 vs. motor de decision en OE3), generando trabajo duplicado o dependencias bloqueantes.",
     "Medio",
     "Retrasos y friccion en la integracion final entre el modulo de pronostico y el motor de decision.",
     "Media",
     "Definir un contrato de formato de salida claro entre el modulo de pronostico y el motor de decision (columnas fecha_hora/prediccion/cuantiles de incertidumbre), para que ambos avancen en paralelo sin bloquearse mutuamente."),

    ("Financiero",
     "Variacion en los costos de servicios en la nube o de APIs de inteligencia artificial usadas durante el desarrollo (riesgo reconocido en el Anexo 1, 6.4.7).",
     "Medio",
     "Podria exceder el presupuesto asignado al rubro de tecnologia si se requiere escalar el computo en la nube.",
     "Baja",
     "Priorizar computo local (estrategia seguida durante todo el desarrollo) y reservar servicios en la nube solo si el hardware disponible resulta insuficiente para una tarea puntual."),

    ("Social",
     "Los usuarios de prueba (muestra de 3-5 personas para OE4) no comprenden con facilidad conceptos de incertidumbre/probabilidad presentados en el dashboard.",
     "Medio",
     "Compromete la validez de la validacion funcional planeada para OE4 y la utilidad percibida de la herramienta.",
     "Media",
     "Disenar las imagenes de apoyo a la decision (semaforo, bandas de incertidumbre) con lenguaje simple y visual, validadas de forma iterativa antes de la prueba formal con usuarios."),

    ("Corrupcion del alcance",
     "Tendencia a perseguir mejoras marginales de precision del modelo (nuevas variables, arquitecturas, papers) en vez de cerrar los objetivos pendientes de mayor peso en el cronograma (OE3, OE4), especialmente bajo la fecha limite de Fase 3.",
     "Alto",
     "Consumo de tiempo y computo en experimentos que no cambian la recomendacion final, retrasando el motor de decision y su validacion.",
     "Media",
     "Fijar un criterio explicito de cuando detener la busqueda de mejoras del modelo (p. ej. significancia estadistica no alcanzada tras N intentos) y pasar a OE3; documentar los intentos descartados en vez de seguir intentando variantes."),

    ("Social / Etico",
     "El usuario final trata la recomendacion del motor de decision como una garantia de resultado economico, en vez de un apoyo a su propio criterio -- justo la restriccion que el Anexo 1 exige evitar (6.3.3).",
     "Alto",
     "Perdidas economicas del usuario atribuidas erroneamente a la herramienta, y perdida de confianza en la plataforma.",
     "Baja",
     "Presentar siempre la incertidumbre junto con la recomendacion puntual; incluir advertencias explicitas de que el sistema no ejecuta transacciones ni sustituye el criterio profesional del usuario."),

    ("Administrativo",
     "Con el volumen de notebooks, scripts y resultados acumulados durante el desarrollo, no alcanzar a consolidar todo en la memoria IEEE, el manual tecnico/de usuario, el poster y el video de la semana 16.",
     "Alto",
     "Entregables incompletos o de baja calidad en la recta final, pese a que el trabajo tecnico si se haya completado.",
     "Media",
     "Empezar la consolidacion de resultados en un documento formal desde ya (no esperar a la semana 16); mantener el README y los resultados de cada fase ya redactados como insumo directo para la memoria."),

]

COLUMNAS = ["Tipo de Riesgo", "Descripción", "Impacto", "Descripción del Impacto", "Probabilidad", "Medida de mitigación"]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Analisis de Riesgos"

azul = "1F3864"
relleno_header = PatternFill(start_color=azul, end_color=azul, fill_type="solid")
fuente_header = Font(color="FFFFFF", bold=True, size=11)
borde = Border(*(Side(style="thin", color="B7B7B7"),) * 4)
alineacion = Alignment(wrap_text=True, vertical="top", horizontal="left")

for col_idx, encabezado in enumerate(COLUMNAS, start=1):
    celda = ws.cell(row=1, column=col_idx, value=encabezado)
    celda.fill = relleno_header
    celda.font = fuente_header
    celda.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    celda.border = borde

for row_idx, fila in enumerate(FILAS, start=2):
    for col_idx, valor in enumerate(fila, start=1):
        celda = ws.cell(row=row_idx, column=col_idx, value=valor)
        celda.alignment = alineacion
        celda.border = borde
        if col_idx == 3:  # Impacto
            celda.font = Font(bold=True)
        if col_idx == 5:  # Probabilidad
            celda.font = Font(bold=True)

anchos = [18, 45, 10, 45, 12, 45]
for i, ancho in enumerate(anchos, start=1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = ancho

for row_idx in range(2, len(FILAS) + 2):
    ws.row_dimensions[row_idx].height = 75
ws.row_dimensions[1].height = 30
ws.freeze_panes = "A2"

ruta = RAIZ / "docs" / "tabla_riesgos_proyecto.xlsx"
ruta.parent.mkdir(exist_ok=True)
wb.save(ruta)
print("Guardado:", ruta)
print(f"{len(FILAS)} riesgos, {len(COLUMNAS)} columnas")
