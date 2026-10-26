"""Genera docs/Referencias_Augur_paper.docx: todo lo que el proyecto leyó o usó (papers, regulación, datos de XM,
software y herramientas), con su cita en formato IEEE, el enlace y para qué se usó en Augur.

Marcas de estado de cada entrada:
  C = citada en el informe final;  L = leída o usada en el proyecto, todavía no citada;
  S = sugerida para respaldar un método que el proyecto sí usa (no se leyó el original);
  V = verificar los datos bibliográficos contra el original antes de citar.
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = Path(__file__).resolve().parents[1]
DESTINO = RAIZ / "docs" / "Referencias_Augur_paper.docx"
AMBAR = RGBColor(0xB4, 0x53, 0x09)
GRIS = RGBColor(0x55, 0x55, 0x55)

# (estado, cita IEEE, enlace o DOI, uso en Augur)
SECCIONES = [
    ("1. Pronóstico de precios de electricidad: métodos que usa el modelo", [
        ("C", "R. Weron, \"Electricity price forecasting: A review of the state-of-the-art with a look into the future,\" "
              "Int. J. Forecast., vol. 30, no. 4, pp. 1030–1081, 2014.",
         "doi: 10.1016/j.ijforecast.2014.08.008",
         "En el informe: las series de precio eléctrico tienen saltos, estacionalidad múltiple y cambios de régimen, y las tres "
         "familias de métodos de pronóstico (estadísticos, aprendizaje automático y redes profundas)."),
        ("C", "J. Lago, G. Marcjasz, B. De Schutter y R. Weron, \"Forecasting day-ahead electricity prices: A review of "
              "state-of-the-art algorithms, best practices and an open-access benchmark,\" Appl. Energy, vol. 293, "
              "art. 116983, 2021.",
         "doi: 10.1016/j.apenergy.2021.116983 · arXiv:2008.08004",
         "Define el LEAR (un LASSO por hora) y las buenas prácticas de evaluación: Diebold-Mariano, rMAE contra "
         "persistencia, protocolo day-ahead. En el informe fija el criterio de precisión (rMAE < 1 frente a la persistencia, "
         "Tabla de requerimientos) y es el referente de buenas prácticas. Respalda el rMAE mensual y la persistencia de 7 días."),
        ("C", "F. Ziel y R. Weron, \"Day-ahead electricity price forecasting with high-dimensional structures: Univariate "
              "vs. multivariate modeling frameworks,\" Energy Econ., vol. 70, pp. 396–420, 2018.",
         "doi: 10.1016/j.eneco.2017.12.016",
         "Explica el salto de medianoche: las 24 horas de un día salen de una sola subasta (bitácora del 18-sep). En el informe "
         "respalda modelar el día como un vector de 24 horas (marcos multivariados)."),
        ("C", "J. Nowotarski y R. Weron, \"Computing electricity spot price prediction intervals using quantile regression "
              "and forecast averaging,\" Comput. Stat., vol. 30, no. 3, pp. 791–803, 2015.",
         "doi: 10.1007/s00180-014-0523-0",
         "QRA: combinador del modelo en producción (pesos no negativos e intercepto por franja de 6 h para el predespacho y "
         "el promedio de los tres CatBoost, objetivo MAE) y del ensamble anterior de seis votantes. En el informe también "
         "respalda el paso al pronóstico probabilístico."),
        ("C", "L. Prokhorenkova, G. Gusev, A. Vorobev, A. V. Dorogush y A. Gulin, \"CatBoost: unbiased boosting with "
              "categorical features,\" en Adv. Neural Inf. Process. Syst. 31 (NeurIPS 2018), 2018.",
         "arXiv:1706.09516",
         "Algoritmo de los tres modelos de producción, que corrigen el predespacho: pronostican log(precio/predespacho) con "
         "pérdida MAE (1.500 árboles, profundidad 6, tasa 0,05)."),
        ("C", "T. Chen y C. Guestrin, \"XGBoost: A scalable tree boosting system,\" en Proc. 22nd ACM SIGKDD, 2016, "
              "pp. 785–794.",
         "doi: 10.1145/2939672.2939785",
         "Votante del ensamble de la primera etapa (notebook 06) y uno de los modelos del benchmark del escenario realista. En el "
         "informe, Solución 1 (ensamble de seis modelos)."),
        ("S", "G. Ke, Q. Meng, T. Finley, T. Wang, W. Chen, W. Ma, Q. Ye y T.-Y. Liu, \"LightGBM: A highly efficient "
              "gradient boosting decision tree,\" en Adv. Neural Inf. Process. Syst. 30 (NIPS 2017), 2017.",
         "proceedings.neurips.cc (NIPS 2017)",
         "LightGBM en el benchmark (empata con CatBoost) y base del HistGradientBoosting del detector de picos de 72 h."),
        ("S", "L. Breiman, \"Random forests,\" Mach. Learn., vol. 45, no. 1, pp. 5–32, 2001.",
         "doi: 10.1023/A:1010933404324",
         "Random Forest en el benchmark del escenario realista (realista_benchmark_modelos.py) y origen de la importancia por "
         "permutación del notebook 06."),
        ("C", "R. Tibshirani, \"Regression shrinkage and selection via the lasso,\" J. R. Stat. Soc. Ser. B, vol. 58, "
              "no. 1, pp. 267–288, 1996.",
         "doi: 10.1111/j.2517-6161.1996.tb02080.x",
         "En el informe: la regularización LASSO en la revisión del estado del arte. En el código: el LEAR (lear_24h.py) y el "
         "Lasso del benchmark, que pierde en 3 de 6 periodos."),
        ("C", "K. G. Olivares, C. Challu, G. Marcjasz, R. Weron y A. Dubrawski, \"Neural basis expansion analysis with "
              "exogenous variables: Forecasting electricity prices with NBEATSx,\" Int. J. Forecast., vol. 39, no. 2, "
              "pp. 884–900, 2023.",
         "doi: 10.1016/j.ijforecast.2022.03.001",
         "N-BEATSx, votante del ensamble de la primera etapa (notebook 09). En el informe, Solución 1."),
        ("C", "C. Challu, K. G. Olivares, B. N. Oreshkin, F. Garza Ramírez, M. Mergenthaler Canseco y A. Dubrawski, "
              "\"NHITS: Neural hierarchical interpolation for time series forecasting,\" en Proc. AAAI, vol. 37, no. 6, "
              "2023, pp. 6989–6997.",
         "doi: 10.1609/aaai.v37i6.25854",
         "N-HiTS, votante del ensamble de la primera etapa (notebook 09). En el informe, Solución 1."),
        ("S", "B. N. Oreshkin, D. Carpov, N. Chapados y Y. Bengio, \"N-BEATS: Neural basis expansion analysis for "
              "interpretable time series forecasting,\" en Proc. ICLR, 2020.",
         "arXiv:1905.10437",
         "Arquitectura original de la que derivan N-BEATSx y N-HiTS."),
        ("C", "T. Bollerslev, \"Generalized autoregressive conditional heteroskedasticity,\" J. Econometrics, vol. 31, "
              "no. 3, pp. 307–327, 1986.",
         "doi: 10.1016/0304-4076(86)90063-1",
         "ARX+GARCH(1,1) (notebook 08): varianza condicional del error y bandas propias; en el informe, Solución 1 (ensamble de seis modelos)."),
        ("C", "D. B. Nelson, \"Conditional heteroskedasticity in asset returns: A new approach,\" Econometrica, vol. 59, "
              "no. 2, pp. 347–370, 1991.",
         "doi: 10.2307/2938260",
         "Introduce el EGARCH y la distribución de error generalizada (GED): base del votante GARCH-GED del ensamble de la "
         "primera etapa (en el informe, Solución 1) y de las variantes EGARCH-t y EGARCH-skewt probadas después."),
        ("S", "S. J. Taylor y B. Letham, \"Forecasting at scale,\" Am. Stat., vol. 72, no. 1, pp. 37–45, 2018.",
         "doi: 10.1080/00031305.2017.1380080",
         "Prophet (notebook 04), el modelo más débil; descartado."),
        ("L", "B. Uniejewski, R. Weron y F. Ziel, \"Variance stabilizing transformations for electricity spot price "
              "forecasting,\" IEEE Trans. Power Syst., vol. 33, no. 2, pp. 2219–2229, 2018.",
         "doi: 10.1109/TPWRS.2017.2734563",
         "Transformación asinh. Probada: empeora 6-16 % en Colombia (los picos son cambios de nivel sostenidos)."),
        ("L", "B. Uniejewski, \"Variance stabilizing transformations for electricity price forecasting in periods of increased "
              "volatility,\" Electr. Power Syst. Res., vol. 257, art. 112992, 2026.",
         "doi: 10.1016/j.epsr.2026.112992",
         "asinh parametrizado con promediado AVGroll sobre el Ridge directo: mejora de solo 0,7 %. Su combinación por hora "
         "del día inspiró aplicar el combinador por franja."),
        ("L", "G. Dutot, M. Zaffran, O. Féron y Y. Goude, \"Adaptive probabilistic forecasting of French electricity spot "
              "prices,\" preprint, 2024.",
         "arXiv:2405.15359",
         "Agregación adaptativa de expertos (BOA). Probada con pesos móviles de 30-60 días: peor que la ventana "
         "expansiva (p < 0,01)."),
    ]),
    ("2. Evaluación, validación e incertidumbre", [
        ("C", "F. X. Diebold y R. S. Mariano, \"Comparing predictive accuracy,\" J. Bus. Econ. Stat., vol. 13, no. 3, "
              "pp. 253–263, 1995.",
         "doi: 10.1080/07350015.1995.10524599",
         "Prueba DM: en el informe, criterio de aceptación (p < 0,05) y comparación con los antecedentes; en los experimentos, "
         "para aceptar o rechazar cada cambio (benchmark, walk-forward, curva50, detector de 72 h)."),
        ("C", "W. K. Newey y K. D. West, \"A simple, positive semi-definite, heteroskedasticity and autocorrelation "
              "consistent covariance matrix,\" Econometrica, vol. 55, no. 3, pp. 703–708, 1987.",
         "doi: 10.2307/1913610",
         "Varianza HAC de Newey-West en la prueba DM, porque los errores horarios están autocorrelacionados."),
        ("C", "L. J. Tashman, \"Out-of-sample tests of forecasting accuracy: An analysis and review,\" Int. J. Forecast., "
              "vol. 16, no. 4, pp. 437–450, 2000.",
         "doi: 10.1016/S0169-2070(00)00065-0",
         "Evaluación fuera de muestra con origen móvil (walk-forward): criterio de robustez del informe (rMAE < 1 en todos los "
         "meses) y los 6 orígenes ENSO y el walk-forward realista."),
        ("C", "R. J. Hyndman y A. B. Koehler, \"Another look at measures of forecast accuracy,\" Int. J. Forecast., "
              "vol. 22, no. 4, pp. 679–688, 2006.",
         "doi: 10.1016/j.ijforecast.2006.03.001",
         "Métricas adimensionales (MASE, error relativo a la persistencia) del criterio de precisión del informe."),
        ("C", "K. Maciejowska, A. Lipiecki y B. Uniejewski, \"Statistical and economic evaluation of forecasts in "
              "electricity markets: Beyond RMSE and MAE,\" Energy Convers. Manag., vol. 356, art. 121408, 2026.",
         "doi: 10.1016/j.enconman.2026.121408 · arXiv:2511.13616",
         "En el informe: un menor error no garantiza mejores decisiones de mercado, y la traducción a pesos no sustituye una "
         "evaluación económica completa. En el código: métricas de forma del día (Corr-f, MHD, MPD; metricas_forma_2026.py)."),
        ("C", "Y. Romano, E. Patterson y E. Candès, \"Conformalized quantile regression,\" en Adv. Neural Inf. Process. "
              "Syst. 32 (NeurIPS 2019), 2019.",
         "arXiv:1905.03222",
         "Inferencia conforme de las bandas [q10, q90]: q10 y q90 = q50·exp(∓m), con m el cuantil conforme de |log(y/q50)| de "
         "la misma hora en los 60 días anteriores ya publicados. También fija el criterio de cobertura (80 ± 5 %)."),
        ("C", "V. Vovk, A. Gammerman y G. Shafer, Algorithmic Learning in a Random World, 2.ª ed. Cham, Suiza: Springer, 2022.",
         "doi: 10.1007/978-3-031-06649-8",
         "Fundamento de la predicción conforme; en el informe, junto a Romano et al., para las bandas por grupos (hora)."),
        ("S", "I. Gibbs y E. Candès, \"Adaptive conformal inference under distribution shift,\" en Adv. Neural Inf. "
              "Process. Syst. 34 (NeurIPS 2021), 2021.",
         "arXiv:2106.00170",
         "Respaldo de la calibración conforme con ventana móvil de 60 días, que se adapta a cambios de distribución."),
        ("L", "T. Gneiting y A. E. Raftery, \"Strictly proper scoring rules, prediction, and estimation,\" J. Am. Stat. "
              "Assoc., vol. 102, no. 477, pp. 359–378, 2007.",
         "doi: 10.1198/016214506000001437",
         "Interval Score (Winkler) para comparar bandas."),
        ("S", "G. W. Brier, \"Verification of forecasts expressed in terms of probability,\" Mon. Weather Rev., vol. 78, "
              "no. 1, pp. 1–3, 1950.",
         "doi: 10.1175/1520-0493(1950)078<0001:VOFEIT>2.0.CO;2",
         "Puntaje de Brier y BSS del detector de picos y desplomes."),
        ("L", "A. H. Murphy, \"Skill scores based on the mean square error and their relationships to the correlation "
              "coefficient,\" Mon. Weather Rev., vol. 116, no. 12, pp. 2417–2424, 1988.",
         "doi: 10.1175/1520-0493(1988)116<2417:SSBOTM>2.0.CO;2",
         "Descomposición del error por día (nivel, amplitud, fase) del diagnóstico de amortiguamiento."),
        ("S", "H. Hewamalage, K. Ackermann y C. Bergmeir, \"Forecast evaluation for data scientists: Common pitfalls and "
              "best practices,\" Data Min. Knowl. Discov., vol. 37, no. 2, pp. 788–832, 2023.",
         "doi: 10.1007/s10618-022-00894-5",
         "Conocido a través de Albani et al. (no se leyó el original): respalda que los modelos complejos no siempre ganan de "
         "forma significativa."),
        ("C", "S. M. Lundberg y S.-I. Lee, \"A unified approach to interpreting model predictions,\" en Adv. Neural Inf. "
              "Process. Syst. 30 (NIPS 2017), 2017.",
         "arXiv:1705.07874",
         "SHAP sobre XGBoost (notebook 06): coincide con la importancia por permutación en las tres variables principales "
         "(Spearman 0,811)."),
        ("S", "J. H. Friedman, \"Greedy function approximation: A gradient boosting machine,\" Ann. Stat., vol. 29, no. 5, "
              "pp. 1189–1232, 2001.",
         "doi: 10.1214/aos/1013203451",
         "Gráficas de dependencia parcial (PDP) del notebook 06 y base del gradient boosting."),
        ("C", "J. W. Tukey, Exploratory Data Analysis. Reading, MA: Addison-Wesley, 1977.",
         "—",
         "Rango intercuartílico: justifica los percentiles 25 y 75 como umbrales por defecto del motor."),
        ("S", "C. W. J. Granger, \"Investigating causal relations by econometric models and cross-spectral methods,\" "
              "Econometrica, vol. 37, no. 3, pp. 424–438, 1969.",
         "doi: 10.2307/1912791",
         "Causalidad de Granger de aportes y embalse sobre el precio (granger_hidrologia.py)."),
        ("L", "T. Emre, \"The influence of hydroelectric power generation and water level variability on market clearing "
              "prices: A case study of Türkiye,\" Int. J. Electr. Power Energy Syst., vol. 170, art. 110865, 2025.",
         "doi: 10.1016/j.ijepes.2025.110865",
         "Metodología que se siguió para la prueba de Granger sobre la serie diaria."),
    ]),
    ("3. Cambio de régimen, puntos de cambio y decisión", [
        ("L", "J. Nasiadka, W. Nitka y R. Weron, \"Calibration window selection based on change-point detection for "
              "forecasting electricity prices,\" en Computational Science – ICCS 2022, LNCS, Springer, 2022, pp. 278–284.",
         "doi: 10.1007/978-3-031-08757-8_24 · arXiv:2204.00872",
         "Selección de ventana por puntos de cambio, implementada en ventana_calibracion_not_24h.py."),
        ("L", "R. Baranowski, Y. Chen y P. Fryzlewicz, \"Narrowest-over-threshold detection of multiple change-points and "
              "change-point-like features,\" J. R. Stat. Soc. Ser. B, vol. 81, no. 3, pp. 649–672, 2019.",
         "doi: 10.1111/rssb.12322 · paquete R: cran.r-project.org/package=not",
         "Detector NOT reimplementado en Python (not_changepoint.py) y validado con series sintéticas."),
        ("L", "R. Huisman y R. Mahieu, \"Regime jumps in electricity prices,\" Energy Econ., vol. 25, no. 5, "
              "pp. 425–434, 2003.",
         "doi: 10.1016/S0140-9883(03)00041-0",
         "Modelo de tres regímenes; origen de la línea de regímenes en EPF."),
        ("L", "A. Das, S. Schlüter y L. Schneider, \"Regime-aware conditional neural processes with multi-criteria decision support for "
              "operational electricity price forecasting,\" Energy Econ., vol. 157, art. 109233, 2026.",
         "doi: 10.1016/j.eneco.2026.109233 · arXiv:2508.00040",
         "Regímenes con HMM en el mercado alemán; confirma que un régimen minoritario (como El Niño) es el caso difícil."),
        ("C", "J. Gama, I. Žliobaitė, A. Bifet, M. Pechenizkiy y A. Bouchachia, \"A survey on concept drift adaptation,\" "
              "ACM Comput. Surv., vol. 46, no. 4, art. 44, 2014.",
         "doi: 10.1145/2523813",
         "Cambio de distribución: respalda el método «rodante» (percentiles de una ventana móvil causal de 30 días sobre q50), "
         "que se adapta al régimen vigente."),
        ("C", "A. Wald, Statistical Decision Functions. Nueva York: Wiley, 1950.",
         "—",
         "Criterio maximin: el método activo es el de mayor ventaja en la mitad más débil del periodo, no el de mayor promedio."),
        ("C", "A. Garivier y E. Moulines, \"On upper-confidence bound policies for switching bandit problems,\" en "
              "Algorithmic Learning Theory (ALT 2011), LNCS, vol. 6925, Springer, 2011, pp. 174–188.",
         "doi: 10.1007/978-3-642-24412-4_16",
         "Trabajo futuro del informe: evaluar el criterio de estabilidad del motor con una formulación de bandits no "
         "estacionarios."),
        ("L", "G. A. Singh, \"Regime-aware deep learning for financial forecasting: an adaptive short-term bias "
              "compensation framework for the Warsaw Stock Exchange,\" Expert Syst. Appl., vol. 332, art. 133584, 2027.",
         "doi: 10.1016/j.eswa.2026.133584",
         "STBC por régimen (fuera de energía). Probado fielmente: λ = 0 en validación; no aporta."),
        ("S", "R. I. Chang, C. H. Wang, L. C. Wei y Y. F. Lu, \"LSTM with short-term bias compensation to determine "
              "trading strategy under black swan events of Taiwan ETF50 stock,\" Appl. Sci., vol. 14, no. 18, art. 8576, 2024.",
         "doi: 10.3390/app14188576",
         "STBC original (λ fijo); conocido a través de Singh (no se leyó el original)."),
        ("L", "V. Le Guen y N. Thome, \"Shape and time distortion loss for training deep time series forecasting "
              "models,\" en Adv. Neural Inf. Process. Syst. 32 (NeurIPS 2019), 2019.",
         "arXiv:1909.09020",
         "Pérdida de forma DILATE, probada contra el amortiguamiento en las rampas: recupera la amplitud pero pasa el error a "
         "la forma (49 % del error cuadrático), así que no se adoptó."),
        ("S", "R. J. Hyndman y M. S. Ullah, \"Robust forecasting of mortality and fertility rates: A functional data "
              "approach,\" Comput. Stat. Data Anal., vol. 51, no. 10, pp. 4942–4956, 2007.",
         "doi: 10.1016/j.csda.2006.07.028",
         "Método del análisis funcional de la curva diaria (fpca_hyndman_24h.py, vía Gallón y Barrientos): curva media más "
         "componentes principales funcionales cuyos pesos se pronostican. Probado como votante: peso cero, pero empata al "
         "ensamble en rampa fuerte."),
    ]),
    ("4. Procesamiento de señales (enfoque de Ingeniería Electrónica)", [
        ("C", "A. Savitzky y M. J. E. Golay, \"Smoothing and differentiation of data by simplified least squares "
              "procedures,\" Anal. Chem., vol. 36, no. 8, pp. 1627–1639, 1964.",
         "doi: 10.1021/ac60214a047",
         "Tendencia de la caracterización (OE1): polinomio de orden 3 en ventanas de 721 h. Validado frente a Butterworth, "
         "Kalman, wavelet y promedio móvil en comparar_filtros_tendencia.py."),
        ("C", "A. V. Oppenheim y R. W. Schafer, Discrete-Time Signal Processing, 3.ª ed. Upper Saddle River, NJ: Pearson, "
              "2010.",
         "—",
         "Periodograma de la serie sin tendencia (Ecuación 6 del informe): identifica los periodos dominantes (fs = 1 muestra/h)."),
        ("S", "S. Butterworth, \"On the theory of filter amplifiers,\" Wireless Engineer, vol. 7, pp. 536–541, 1930.",
         "—",
         "Butterworth de fase cero y causal en la comparación de filtros (comparar_filtros_tendencia.py)."),
        ("S", "R. E. Kalman, \"A new approach to linear filtering and prediction problems,\" J. Basic Eng., vol. 82, "
              "no. 1, pp. 35–45, 1960.",
         "doi: 10.1115/1.3662552",
         "Filtro de Kalman en la misma comparación."),
        ("S", "H. E. Rauch, F. Tung y C. T. Striebel, \"Maximum likelihood estimates of linear dynamic systems,\" "
              "AIAA J., vol. 3, no. 8, pp. 1445–1450, 1965.",
         "doi: 10.2514/3.3166",
         "Suavizador RTS (Kalman fuera de línea) de la comparación."),
        ("S", "S. G. Mallat, \"A theory for multiresolution signal decomposition: The wavelet representation,\" IEEE "
              "Trans. Pattern Anal. Mach. Intell., vol. 11, no. 7, pp. 674–693, 1989.",
         "doi: 10.1109/34.192463",
         "Tendencia por wavelet (sym8) de la comparación."),
    ]),
    ("5. Mercados hidro-dominados (comparación internacional)", [
        ("C", "G. Kapoor y N. Wichitaksorn, \"Electricity price forecasting in New Zealand: A comparative analysis of "
              "statistical and machine learning models with feature selection,\" Appl. Energy, vol. 347, art. 121446, "
              "2023.",
         "doi: 10.1016/j.apenergy.2023.121446",
         "En el informe: ninguno de sus 33 modelos baja de MASE 1 (mejor 1,26), mientras Augur supera a su referencia ingenua "
         "en todos los años; se presenta como «superar al ingenuo en una tarea más difícil», no como comparación directa "
         "(transforman con Box-Cox y pronostican precio diario)."),
        ("L", "G. Kapoor, N. Wichitaksorn y W. Zhang, \"Analyzing and forecasting electricity price using regime-switching "
              "models: The case of New Zealand market,\" J. Forecast., vol. 42, no. 8, pp. 2011–2026, 2023.",
         "doi: 10.1002/for.3004",
         "Regímenes de Markov y EVT; más regímenes no mejoran el pronóstico."),
        ("L", "M. T. D. Phan, T. T. Truong, H. P. Ha y D. T. Nguyen, \"Electricity price forecasting across Norway's five "
              "bidding zones in the post-crisis era,\" preprint, 2026.",
         "arXiv:2604.26634",
         "Referencia noruega con LightGBM; comparación de sMAPE y R²."),
        ("L", "A. Brigatto, A. Street y J. Dias Garcia, \"How optimistic inflow forecasts distort dispatch, prices, and "
              "contracts in hydro-dominated power systems: evidence from Brazil,\" preprint, 2026.",
         "arXiv:2607.00504",
         "Idea de «memoria hidrológica» (anomalía acumulada de aportes). Probada: no significativa (p = 0,52)."),
        ("L", "M. B. B. Dias, G. R. S. Lira y V. M. E. Freire, \"Methodology for multi-step forecasting of electricity spot "
              "prices based on neural networks applied to the Brazilian energy market,\" Energies, vol. 17, no. 8, art. 1864, "
              "2024.",
         "doi: 10.3390/en17081864",
         "Poda por multicolinealidad: probada, no funciona y es riesgosa si se implementa literal; referencia brasileña de "
         "pronóstico multipaso."),
        ("L", "E. A. Nunes, V. H. Ferreira y A. da Costa Pinho, \"Information theoretic learning applied to daily streamflow "
              "forecast and its impact on the Brazilian hourly energy spot prices,\" J. Control Autom. Electr. Syst., vol. 35, "
              "no. 5, pp. 949–959, 2024.",
         "doi: 10.1007/s40313-024-01110-z",
         "Pérdida MCC (correntropía), probada en mcc_perdida_robusta_24h.py."),
        ("L", "V. V. L. Albani, R. T. Marcavillaca, P. S. E. Moreira et al., \"Short-term forecasting of forward prices in "
              "the Brazilian electricity market with a hybrid stochastic-neural network model,\" Energy Econ., vol. 148, "
              "art. 108651, 2025.",
         "doi: 10.1016/j.eneco.2025.108651",
         "Compuerta por régimen hidrológico; probada en qra_por_regimen.py: no sirve."),
    ]),
    ("6. Antecedentes en Colombia", [
        ("C", "A. P. Agudelo, J. M. López-Lezama y E. Velilla, \"Predicción del precio de la electricidad en la bolsa "
              "mediante un modelo neuronal no-lineal autorregresivo con entradas exógenas,\" Inf. Tecnol., vol. 26, "
              "no. 6, pp. 99–108, 2015.",
         "doi: 10.4067/S0718-07642015000600012",
         "Antecedente colombiano (Tabla 4 del informe): red NARX con entradas exógenas, datos previos a 2015, sin incertidumbre."),
        ("C", "A. Muñoz-Santiago, J. Urquijo-Vanstrahlengs, A. Castro-Otero y J. Lombana, \"Pronóstico del precio de la "
              "energía en Colombia utilizando modelos ARIMA con IGARCH,\" Rev. Econ. Rosario, vol. 20, no. 1, "
              "pp. 127–161, 2017.",
         "doi: 10.12804/revistas.urosario.edu.co/economia/a.6152",
         "ARIMA con IGARCH: base del votante ARX+GARCH y antecedente colombiano (Tabla 4 del informe)."),
        ("C", "J. Barrientos Marín, E. Tabares Orozco y E. Velilla, \"Forecasting electricity price in Colombia: A "
              "comparison between neural network, ARMA process and hybrid models,\" Int. J. Energy Econ. Policy, "
              "vol. 8, no. 3, pp. 97–106, 2018.",
         "—", "Antecedente colombiano (Tabla 4): ARMAX, NARX e híbrido, mensual; con la prueba DM no hay diferencia "
         "significativa entre modelos."),
        ("C", "S. Gallón y J. Barrientos, \"Forecasting the Colombian electricity spot price under a functional "
              "approach,\" Int. J. Energy Econ. Policy, vol. 11, no. 2, pp. 67–74, 2021.",
         "doi: 10.32479/ijeep.10607 (impreso en el PDF; en Crossref resuelve a otro artículo, conviene citar también "
         "la URL de econjournals.com)",
         "Antecedente colombiano horario (FPCA): su MAPE de 6,7 % a un día se compara con el 10,0 % de Augur, en un periodo "
         "calmo (2000-2017) y suponiendo conocido el precio del día anterior."),
        ("L", "C. A. Ferraro Rodríguez, \"Análisis de correlaciones del precio de bolsa de energía eléctrica con variables "
              "del sector eléctrico y económicas,\" tesis de maestría, Univ. Nacional de Colombia, Medellín, 2025.",
         "repositorio.unal.edu.co/items/05583774-4f00-4fb1-8cfa-ecf426497670",
         "Respalda corregir el predespacho en vez de modelar desde cero (el precio lo fija la oferta marginal)."),
        ("L", "J. D. Becerra Pulido, \"Pronóstico del precio diario de bolsa en el mercado eléctrico colombiano usando "
              "modelos autorregresivos,\" proyecto de grado, Ing. Electrónica, Univ. Santo Tomás, Bogotá, 2025.",
         "repository.usta.edu.co/items/d8dd82e6-6d1c-418d-915d-f035cef594d9",
         "El trabajo más comparable (mismo programa). Usa el precio del día anterior, que XM no publica a tiempo."),
        ("L", "L. Herrera-Mejía, C. D. Hoyos, D. La Rotta, Y. Mejía y J. Castillo, \"Forecasting day-ahead hydropower bids in "
              "the Colombian electricity market: A two-stage machine learning framework,\" Smart Energy, vol. 21, art. 100221, "
              "2026.",
         "doi: 10.1016/j.segy.2025.100221",
         "Inspiró el clasificador de régimen alto (precio > percentil 80) con regresores por régimen."),
        ("C", "International Energy Agency, \"Colombia 2023 – Energy Policy Review,\" IEA, París, 2023.",
         "https://www.iea.org/reports/colombia-2023", "Matriz mayoritariamente hidroeléctrica: cerca del 72 % de la generación."),
        ("V", "Low Carbon Power, \"Electricity in Colombia,\" 2025.",
         "https://lowcarbonpower.org/region/Colombia",
         "Participación hidroeléctrica (70-74 %). Se citaba en un borrador anterior del informe; el final usa la IEA."),
    ]),
    ("7. Regulación (CREG) y normas", [
        ("C", "CREG, \"Resolución CREG 025 de 1995, por la cual se establece el Código de Redes,\" Bogotá, 1995.",
         "gestornormativo.creg.gov.co",
         "Cierre de ofertas a las 08:00 de D-1: el pronóstico con predespacho sirve para decisiones posteriores (compras en "
         "bolsa, coberturas), no para formular ofertas."),
        ("L", "CREG, Resolución 024 de 1995 (modificada por la 011 de 2010) y Resolución 051 de 2009 (modificada por la "
              "101 028 de 2022): formación del precio de bolsa.",
         "gestornormativo.creg.gov.co", "Despacho ideal ex post, MPO del recurso marginal y ΔI."),
        ("L", "CREG, Documento CREG 901 044 de 2023, \"Metodología para la determinación del precio de bolsa — propuesta "
              "de ajuste temporal\" (soporte del Proyecto de Resolución 701 028 de 2023).",
         "gestornormativo.creg.gov.co/gestor/entorno/docs/originales/Proyecto_Resoluci%C3%B3n_CREG_701_028_2023/",
         "Datos del planteamiento: 67 % de capacidad hidráulica, concentración, índice de Lerner."),
        ("L", "CREG, Proyecto de Resolución 701 049 de 2024 y Documento 901 097 de 2024; SIC, Concepto 183210 de 2024.",
         "gestornormativo.creg.gov.co/gestor/entorno/docs/originales/Proyecto_Resoluci%C3%B3n_CREG_701_049_2024/",
         "Ajuste del MPO hidráulico (no adoptado)."),
        ("L", "CREG, Resoluciones 101 056 y 101 059 de 2024 (precio de oferta ajustado; modifican la CREG 026 de 2014).",
         "gestornormativo.creg.gov.co/gestor/entorno/docs/resolucion_creg_101-59_2024.htm",
         "Cambios de regla dentro del periodo de entrenamiento (sep-oct 2024)."),
        ("L", "CREG, Resolución 101 127 de 2026 y Documento 901 431 de 2026, \"Medidas transitorias para asegurar el "
              "abastecimiento de energía ante el fenómeno de El Niño 2026-2027\".",
         "gestornormativo.creg.gov.co/gestor/entorno/docs/originales/Resoluci%C3%B3n_CREG_101_127_2026/",
         "Cambió la formación del precio desde sep-2026; explica el mayor error de sep-oct 2026."),
        ("L", "CREG, Resolución 157 de 2011 (modificada por la 211 de 2015) y Concepto 6762 de 2020: versiones de "
              "liquidación TX1, TX2, TXR y TXF.",
         "gestornormativo.creg.gov.co/gestor/entorno/docs/ee_2_9_2_3.htm",
         "Por qué el precio real llega con unos 3 días de rezago (corte del escenario realista)."),
        ("C", "CREG, \"Resolución CREG 178 de 2015, por la cual se establecen medidas para garantizar la prestación del "
              "servicio público de energía eléctrica,\" Bogotá, 27 de octubre de 2015.",
         "gestornormativo.creg.gov.co/gestor/entorno/docs/resolucion_creg_0178_2015.htm",
         "El «cambio regulatorio de 2015» del Anexo 1, parte transitoria: durante El Niño 2015-2016 fijó un piso al "
         "precio de escasez y una opción para térmicas a combustibles líquidos, por seis meses (la Res. 052 de 2016 no la "
         "prorrogó). En el informe (Ventana de datos), primera razón para excluir 2015-2016: el mercado operó con reglas "
         "de excepción."),
        ("C", "CREG, \"Resolución CREG 140 de 2017, por la cual se define el precio marginal de escasez del Cargo por "
              "Confiabilidad y se hacen modificaciones a la Resolución CREG 071 de 2006,\" Bogotá, 2017.",
         "gestornormativo.creg.gov.co/gestor/entorno/docs/resolucion_creg_0140_2017.htm",
         "Precio marginal de escasez, vigente desde el 1-dic-2017: cambió la regla que activa la condición crítica y pone "
         "techo al precio que paga la demanda. En el informe (Ventana de datos), segunda razón para empezar la serie en 2019."),
        ("L", "CREG, Resoluciones 071 de 2006 y 101 066 de 2024: precio de escasez y precio de transacciones en bolsa (PTB).",
         "gestornormativo.creg.gov.co/gestor/entorno/docs/resolucion_creg_0071_2006.htm",
         "Limitación: el PTB fue menor que el PBN en 360 horas de 2026."),
        ("C", "CREG, \"Resolución CREG 101 018 de 2022, por la cual se crea el Sistema de Información del Mercado de Energía "
              "Mayorista, SIMEM,\" Bogotá, 2022.", "gestornormativo.creg.gov.co/gestor/entorno/docs/resolucion_creg_101-18_2022.htm",
         "Crea el SIMEM: respalda el uso exclusivo de información pública de XM/SIMEM; el sistema no actúa como agente."),
        ("C", "Congreso de la República de Colombia, \"Ley Estatutaria 1581 de 2012, por la cual se dictan disposiciones "
              "generales para la protección de datos personales,\" Bogotá, Colombia, oct. 2012.", "—",
         "Restricción de diseño: el dashboard no recolecta ni almacena datos personales."),
        ("C", "ISO 8601-1:2019, Date and time — Representations for information interchange — Part 1: Basic rules.", "—",
         "Marcas de tiempo AAAA-MM-DD hh:mm:ss en hora legal de Colombia (UTC−5) para unir fuentes horarias, diarias y mensuales."),
        ("C", "ISO/IEC 25010:2011, Systems and software engineering — Systems and software Quality Requirements and "
              "Evaluation (SQuaRE) — System and software quality models.", "—",
         "Mantenibilidad: el motor lee el pronóstico por un contrato de datos (fuentes_pronostico.json) y no depende del modelo."),
        ("C", "ACIEM y Consejo Profesional Nacional de Ingenierías Eléctrica, Mecánica y Profesiones Afines, Manual de "
              "Referencia de Tarifas en Ingeniería. Cundinamarca, Colombia, 2015.", "—",
         "Honorarios del presupuesto del proyecto (Anexo 1)."),
    ]),
    ("8. Datos: XM y otras fuentes", [
        ("C", "XM S.A. E.S.P., \"API de datos del Sistema de Información del Mercado de Energía Mayorista\" "
              "(servapibi.xm.com.co).",
         "https://servapibi.xm.com.co · librería pydataxm: github.com/EquipoAnaliticaXM/API_XM",
         "Fuente principal. Métricas usadas: PrecBolsNaci (precio de bolsa), DemaReal y DemaCome (demanda), Gene por "
         "recurso y GeneIdea (generación), AporEner y AporEnerMediHist (aportes), VoluUtilDiarEner, PorcVoluUtilDiar y "
         "CapaUtilDiarEner (embalse), PrecOferDesp y MaxPrecOferNal (ofertas), PrecEscaMarg y PrecEsca (escasez), "
         "PrecTransBolsa, RestAliv, ImpoEner, ExpoEner, PerdidasEner; catálogos ListadoMetricas y ListadoRecursos."),
        ("C", "XM S.A. E.S.P., \"Informes de despacho: predespacho ideal y despacho programado,\" biblioteca pública del "
              "portal xm.com.co.",
         "https://www.xm.com.co/generacion/informes-despacho · archivos iMAR/dMAR vía "
         "api-portalxm.xm.com.co/administracion-archivos",
         "Insumo central del modelo en producción: el predespacho ideal que XM publica el día anterior."),
        ("L", "XM S.A. E.S.P., ofertas de precio (archivos OFEI).", "Biblioteca del portal XM",
         "Probadas como variable: los precios no se publican a tiempo (sería fuga de información)."),
        ("L", "XM S.A. E.S.P., Sinergox: tablero del Precio de Bolsa Nacional.", "https://sinergox.xm.com.co",
         "Escala de color del mapa de calor por precio esperado."),
        ("L", "XM S.A. E.S.P., \"Resolución CREG 101 127 de 2026 — condición del sistema\".",
         "https://www.xm.com.co/resoluciones/operacion-y-mercado/resolucion-creg-101-127-de-2026-condicion-del-sistema",
         "Seguimiento de la activación de las medidas de El Niño."),
        ("C", "NOAA Climate Prediction Center, \"Cold & Warm Episodes by Season — Oceanic Niño Index (ONI)\".",
         "https://www.cpc.ncep.noaa.gov/data/indices/oni.ascii.txt",
         "Índice ONI (El Niño / La Niña) como variable y para definir los orígenes del walk-forward."),
        ("L", "Paquete holidays (Python): festivos de Colombia.", "https://github.com/vacanza/holidays",
         "Variable de calendario (festivos)."),
    ]),
    ("9. Referencias del Anexo 1 (formulación) que no usa el informe final", [
        ("L", "Congreso de Colombia, \"Ley 23 de 1982, sobre derechos de autor,\" Diario Oficial n.º 35.949, 19 de febrero "
              "de 1982.",
         "normas.cra.gov.co/gestor/docs/ley_0023_1982.htm",
         "Anexo 1, restricciones del diseño: propiedad intelectual del código, los datos y los materiales de terceros."),
        ("L", "Comunidad Andina, \"Decisión 351 de 1993: Régimen común sobre derecho de autor y derechos conexos,\" Lima, "
              "1993.",
         "https://www.comunidadandina.org/StaticFiles/DocOf/DEC351.pdf",
         "Anexo 1, restricciones del diseño, junto con la Ley 23 de 1982."),
        ("L", "ABET, \"Criteria for Accrediting Engineering Programs, 2025–2026,\" Engineering Accreditation Commission, "
              "Baltimore, MD, EE. UU.",
         "https://www.abet.org/accreditation/accreditation-criteria/criteria-for-accrediting-engineering-programs-2025-2026/",
         "Resultado SO2 (diseño en ingeniería): ítem 6 del Anexo 1 y estructura del informe final ABET."),
        ("L", "P. Villegas y H. A. Gómez, \"Medidas regulatorias posteriores a los fenómenos de El Niño 2015-2016 y 2023-2024 "
              "en Colombia, una mirada desde la confiabilidad,\" en XX ERIAC, CIGRE, Ciudad del Este, Paraguay, mayo de 2025.",
         "xm.com.co (Paper N.º 48, Comité C5)",
         "Fuente secundaria sobre los cambios al precio de escasez tras El Niño 2015-2016: en ese periodo el precio de "
         "bolsa superó el precio de escasez el 61,4 % de los días (hasta 2.821 COP/kWh), y el PME rige desde el 1-dic-2017. "
         "Cautela: atribuye a la Res. CREG 088 de 2016 un ajuste del precio de escasez, pero esa resolución trata de la "
         "oferta de GLP; no citarla."),
        ("L", "Globalsyde, \"Software para comercializadoras de energía\" (plataforma en la nube), España.",
         "https://globalsyde.com",
         "Anexo 1, justificación (2): herramienta comercial para comercializadoras minoristas españolas (contratación, "
         "tarificación y facturación), que no aplica al mercado mayorista colombiano."),
        ("L", "Enexpa, \"Servicios para comercializadoras de electricidad y gas en España\".",
         "https://www.enexpa.es",
         "Anexo 1, justificación (2), por la misma razón que Globalsyde."),
        ("V", "SIGE Comercialización, ERP para comercializadoras y distribuidoras de electricidad y gas, España.",
         "asest.es/story/sige-erp-electridad-gas/",
         "Anexo 1, justificación (2). Fabricante y página oficial sin confirmar (un directorio lo atribuye a Grupo Audinfor)."),
    ]),
]

# (librería, versión, para qué en Augur, cita sugerida, licencia)
LIBRERIAS = [
    ("pandas", "2.3.3", "Manejo de todas las series y tablas",
     "W. McKinney, \"Data structures for statistical computing in Python,\" Proc. 9th Python in Science Conf., 2010.",
     "BSD-3"),
    ("NumPy", "2.4.4", "Cálculo numérico",
     "C. R. Harris et al., \"Array programming with NumPy,\" Nature, vol. 585, pp. 357–362, 2020.", "BSD-3"),
    ("SciPy", "1.18.0", "Filtros (Savitzky-Golay, Butterworth), estadística, pruebas",
     "P. Virtanen et al., \"SciPy 1.0,\" Nat. Methods, vol. 17, pp. 261–272, 2020.", "BSD-3"),
    ("scikit-learn", "1.9.0", "Lasso, Random Forest, HistGradientBoosting (detector 72 h), validación",
     "F. Pedregosa et al., \"Scikit-learn: Machine learning in Python,\" JMLR, vol. 12, pp. 2825–2830, 2011.",
     "BSD-3"),
    ("statsmodels", "0.15.0", "Regresión cuantílica (QRA), ARIMA, Granger, pruebas",
     "S. Seabold y J. Perktold, \"Statsmodels: Econometric and statistical modeling with Python,\" Proc. 9th "
     "Python in Science Conf., 2010.", "BSD-3"),
    ("CatBoost", "1.2.10", "Modelos de corrección del predespacho en producción", "Prokhorenkova et al. (2018), sección 1",
     "Apache 2.0"),
    ("LightGBM", "4.7.0", "Benchmark (notebook 11 y escenario realista)", "Ke et al. (2017), sección 1", "MIT"),
    ("XGBoost", "2.1.4", "Modelo de la primera etapa y benchmark", "Chen y Guestrin (2016), sección 1", "Apache 2.0"),
    ("arch", "8.0.0", "ARX+GARCH y GARCH-GED",
     "K. Sheppard et al., arch: ARCH models in Python (software), github.com/bashtage/arch.", "NCSA"),
    ("neuralforecast", "3.2.1", "N-BEATSx y N-HiTS",
     "K. G. Olivares, C. Challu, F. Garza y M. Mergenthaler, NeuralForecast (software), Nixtla, 2022.",
     "Apache 2.0"),
    ("PyTorch / Lightning", "2.14.0 / 2.5.6", "Motor de las redes de neuralforecast",
     "A. Paszke et al., \"PyTorch: An imperative style, high-performance deep learning library,\" NeurIPS, 2019.",
     "BSD-3 / Apache 2.0"),
    ("Prophet (+ CmdStan)", "1.3.0", "Modelo Prophet (descartado)", "Taylor y Letham (2018), sección 1", "MIT"),
    ("SHAP", "0.52.0", "Explicabilidad de XGBoost", "Lundberg y Lee (2017), sección 2", "MIT"),
    ("Optuna", "4.9.0", "Búsqueda de hiperparámetros (armónicos XGBoost/CatBoost)",
     "T. Akiba et al., \"Optuna: A next-generation hyperparameter optimization framework,\" KDD, 2019.", "MIT"),
    ("PyWavelets", "1.8.0", "Filtro wavelet de la comparación de filtros",
     "G. R. Lee et al., \"PyWavelets: A Python package for wavelet analysis,\" JOSS, vol. 4, no. 36, p. 1237, 2019.",
     "MIT"),
    ("pydataxm", "0.3.17", "Descarga de la API de XM", "Equipo de Analítica de XM, github.com/EquipoAnaliticaXM/API_XM",
     "MIT"),
    ("holidays", "0.102", "Festivos de Colombia", "github.com/vacanza/holidays", "MIT"),
    ("requests", "2.33.1", "Descarga del predespacho y del ONI", "—", "Apache 2.0"),
    ("Matplotlib / seaborn", "3.10.8 / 0.13.2", "Figuras del informe y de los notebooks",
     "J. D. Hunter, \"Matplotlib: A 2D graphics environment,\" Comput. Sci. Eng., vol. 9, no. 3, pp. 90–95, 2007; "
     "M. L. Waskom, \"seaborn,\" JOSS, vol. 6, no. 60, p. 3021, 2021.", "PSF / BSD-3"),
    ("Streamlit", "1.63.0", "Dashboard (augurpredictor.streamlit.app)", "streamlit.io", "Apache 2.0"),
    ("Plotly", "7.0.0", "Gráficas interactivas del dashboard", "plotly.com/python", "MIT"),
    ("jsPDF", "2.5.1", "Exportar gráficas y reporte del día a PDF desde el dashboard", "github.com/parallax/jsPDF", "MIT"),
    ("python-docx / openpyxl", "1.2.0 / 3.1.5", "Generación de informes Word y tablas Excel", "—", "MIT"),
    ("Jupyter (ipykernel, nbconvert)", "—", "Notebooks 01 a 12", "jupyter.org", "BSD-3"),
]

# (herramienta, para qué)
HERRAMIENTAS = [
    ("Python 3.12", "Lenguaje de todo el proyecto."),
    ("Git y GitHub (github.com/Ziskito/xm-spot-price-predictor)", "Control de versiones y repositorio público."),
    ("GitHub Actions", "Actualización diaria automática del pronóstico (hacia las 12:00, 15:40 y 18:40)."),
    ("Streamlit Community Cloud", "Publicación del dashboard; cada cambio en main se redespliega."),
    ("Playwright (Microsoft)", "Pruebas del dashboard en navegador y grabación de las tomas del video."),
    ("HyperFrames (HeyGen)", "Edición del video explicativo (HTML renderizado a video)."),
    ("whisper.cpp (modelo small de OpenAI Whisper)", "Tiempos por palabra para los subtítulos del video."),
    ("Gemini TTS (Google), voz Charon", "Voz del video explicativo."),
    ("FFmpeg", "Conversión de las tomas del video."),
    ("Canva", "Presentación del proyecto."),
    ("Microsoft Forms", "Encuestas de validación con usuarios (OE4-2)."),
    ("LibreOffice", "Conversión de los informes a PDF."),
    ("Claude Code (Anthropic)", "Asistente de programación usado durante el desarrollo."),
]


def borde_inferior(par, color="B45309"):
    pPr = par._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    e = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "4"), ("w:color", color)):
        e.set(qn(k), v)
    b.append(e)
    pPr.append(b)


def sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:color"), "auto")
    s.set(qn("w:fill"), color)
    tcPr.append(s)


def texto_celda(celda, texto, negrita=False, tam=8.5):
    celda.text = ""
    r = celda.paragraphs[0].add_run(texto)
    r.bold = negrita
    r.font.size = Pt(tam)


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
for lado in ("left_margin", "right_margin"):
    setattr(sec, lado, Cm(2.2))
sec.top_margin = sec.bottom_margin = Cm(2)
estilo = doc.styles["Normal"]
estilo.font.name = "Calibri"
estilo.font.size = Pt(10.5)
for nombre in ("Heading 1", "Heading 2"):
    doc.styles[nombre].font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    doc.styles[nombre].font.name = "Calibri"

t = doc.add_paragraph()
r = t.add_run("Augur: fuentes, datos, software y referencias")
r.bold, r.font.size = True, Pt(20)
s = doc.add_paragraph()
r = s.add_run("Insumo para el paper · todo lo que el proyecto leyó o usó · versión del 23 de octubre de 2026")
r.font.size, r.font.color.rgb = Pt(10), GRIS
borde_inferior(s)

doc.add_paragraph(
    "Este documento reúne, en un solo lugar, los papers, la regulación, las fuentes de datos, las librerías y las "
    "herramientas que se usaron en Augur, la plataforma de pronóstico horario del precio de bolsa colombiano con motor "
    "de decisión. Cada entrada trae la cita en formato IEEE, el enlace o DOI y una línea sobre para qué se usó en el "
    "proyecto, para que se pueda decidir rápido qué citar y dónde.")
doc.add_paragraph("Cada entrada lleva una marca de estado:")
for marca, desc in (("C", "citada en el informe final del proyecto (datos verificados contra Crossref)."),
                    ("L", "leída o usada en el proyecto, todavía no citada en el informe."),
                    ("S", "sugerida: respalda un método que el proyecto sí usa, pero el original no se leyó."),
                    ("V", "verificar los datos bibliográficos contra el original antes de citarla.")):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(f"[{marca}] ")
    r.bold, r.font.color.rgb = True, AMBAR
    p.add_run(desc)
doc.add_paragraph(
    "Los PDF de la mayoría de los papers de las secciones 1, 3, 5 y 6 están en docs/papers/ del repositorio, con una "
    "ficha de cada uno (y las cautelas al citarlo) en docs/papers/REFERENCIAS.md.")

n = 0
for titulo, entradas in SECCIONES:
    doc.add_heading(titulo, level=1)
    for estado, cita, enlace, uso in entradas:
        n += 1
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.1)
        p.paragraph_format.first_line_indent = Cm(-1.1)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(f"[{estado}]\t")
        r.bold, r.font.color.rgb = True, AMBAR
        p.add_run(cita)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(1.1))
        d = doc.add_paragraph()
        d.paragraph_format.left_indent = Cm(1.1)
        d.paragraph_format.space_after = Pt(7)
        r = d.add_run(enlace + "\n")
        r.font.size, r.font.color.rgb = Pt(9), GRIS
        r = d.add_run("Uso en Augur: ")
        r.bold, r.font.size = True, Pt(9.5)
        d.add_run(uso).font.size = Pt(9.5)

doc.add_heading("10. Librerías de código abierto", level=1)
doc.add_paragraph(
    "Versiones fijadas en requirements.txt (modelado), requirements-produccion.txt (actualización diaria) y "
    "dashboard/requirements.txt (app). Para el paper basta citar las que hacen el trabajo central: pandas, NumPy, "
    "SciPy, scikit-learn, statsmodels, CatBoost y Streamlit.")
tabla = doc.add_table(rows=1, cols=5)
tabla.style = "Table Grid"
tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
anchos = (Cm(3.0), Cm(2.0), Cm(4.6), Cm(5.6), Cm(2.0))
for i, enc in enumerate(("Librería", "Versión", "Para qué en Augur", "Cita sugerida", "Licencia")):
    c = tabla.rows[0].cells[i]
    texto_celda(c, enc, negrita=True, tam=9)
    sombrear(c, "F3EBDB")
for fila in LIBRERIAS:
    celdas = tabla.add_row().cells
    for i, valor in enumerate(fila):
        texto_celda(celdas[i], valor, negrita=(i == 0))
for fila in tabla.rows:
    for i, c in enumerate(fila.cells):
        c.width = anchos[i]

doc.add_heading("11. Herramientas e infraestructura", level=1)
for nombre, uso in HERRAMIENTAS:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(nombre + ": ").bold = True
    p.add_run(uso)

doc.add_heading("12. Material propio del proyecto (para citar o anexar)", level=1)
for nombre, uso in (
        ("Repositorio", "github.com/Ziskito/xm-spot-price-predictor (código, datos procesados, notebooks y bitácora en "
                        "el README)."),
        ("Dashboard en línea", "https://augurpredictor.streamlit.app"),
        ("Documento de funcionamiento", "docs/Funcionamiento_Augur.pdf: variables, métricas, modelos, benchmark, "
                                        "walk-forward y detector de 72 h, con sus tablas."),
        ("Manual de usuario", "docs/Manual_Usuario_augur.pdf (MU-AUG-001)."),
        ("Resultados del escenario realista", "data/processed/resultados/escenario_realista/ (benchmark/RESUMEN.md, "
                                              "walkforward, eventos_72h)."),
        ("Video explicativo", "docs/video/demo_final.mp4 (3:53) con subtítulos y guion."),
        ("Formulación", "J. D. Barceló Barraza y R. A. Dede Perdomo, Anexo 1 – Formulación del proyecto, Universidad "
                        "del Norte, 2026.")):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(nombre + ": ").bold = True
    p.add_run(uso)

pie = doc.add_paragraph()
pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = pie.add_run(f"{n} referencias bibliográficas, regulatorias y de datos · {len(LIBRERIAS)} librerías · "
                f"{len(HERRAMIENTAS)} herramientas")
r.font.size, r.font.color.rgb = Pt(9), GRIS

DESTINO.parent.mkdir(exist_ok=True)
doc.save(str(DESTINO))
print("Guardado:", DESTINO, "|", n, "referencias")
